from __future__ import annotations

import re
from datetime import time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError

RESIDUAL_MIX_HEADERS = (
    None,
    "RE Total",
    "RE unspecified",
    "RE biomass",
    "RE solar",
    "RE geothermal",
    "RE wind",
    "RE hydro",
    "Nuclear",
    "FO Total",
    "FO unspecified",
    "FO hard coal",
    "FO lignite",
    "FO oil",
    "FO gas",
    "Untracked %",
    "CO2 (gCO2/kWh)",
    "Rad waste (mg/kWh)",
)
ATTRIBUTE_HEADERS = (
    "Country code",
    "Production mix CO2",
    "Residual mix CO2",
    "Supplier mix CO2",
)
WASTE_HEADERS = (
    "Country code",
    "Production mix RW",
    "Residual mix RW",
    "Supplier mix RW",
)
SHARE_NAMES = (
    "renewable_total",
    "renewable_unspecified",
    "renewable_biomass",
    "renewable_solar",
    "renewable_geothermal",
    "renewable_wind",
    "renewable_hydro",
    "nuclear",
    "fossil_total",
    "fossil_unspecified",
    "fossil_hard_coal",
    "fossil_lignite",
    "fossil_oil",
    "fossil_gas",
)


class AibRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    country_code: str = Field(pattern=r"^[A-Z]{2}$")
    available: bool
    direct_co2_g_per_kwh: Decimal | None = None
    radioactive_waste_mg_per_kwh: Decimal | None = None
    shares: dict[str, Decimal] = Field(default_factory=dict)
    untracked_share: Decimal | None = None
    source_row: int
    residual_mix_row: int
    source_sheet: str = "Residual Mixes"

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet=self.source_sheet,
            source_table="Residual mix CO2",
        )


class AibParser:
    """Parse the versioned country residual-mix contract from AIB's workbook."""

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(self, path: Path) -> tuple[AibRow, ...]:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as error:
            raise PermanentIngestionError("AIB file is not a readable XLSX workbook") from error
        try:
            residual_name = next(
                (
                    name
                    for name in workbook.sheetnames
                    if name.startswith("Residual Mixes")
                    and "Comparison" not in name
                    and "," not in name
                ),
                None,
            )
            if residual_name is None:
                raise PermanentIngestionError(
                    "AIB workbook sheets changed; parser version must be updated"
                )
            residual = workbook[residual_name]
            self._assert_residual_headers(residual)
            rows: list[AibRow] = []
            seen_codes: set[str] = set()
            for row_number, values in enumerate(
                residual.iter_rows(min_row=2, values_only=True), start=2
            ):
                code = self._code(values[0] if values else None)
                if code is None:
                    continue
                if code in seen_codes:
                    raise PermanentIngestionError(f"duplicate AIB country code: {code}")
                seen_codes.add(code)
                shares: dict[str, Decimal] = {}
                for name, value in zip(SHARE_NAMES, values[1:15], strict=True):
                    parsed = self._share(value, row_number)
                    if parsed is not None:
                        shares[name] = parsed
                untracked = self._share(values[15], row_number)
                co2_value = self._decimal(values[16], row_number)
                waste_value = self._decimal(values[17], row_number)
                has_reported_mix = any(value != 0 for value in shares.values()) or (
                    untracked is not None and untracked != 0
                )
                available = co2_value is not None and (co2_value != 0 or has_reported_mix)
                rows.append(
                    AibRow(
                        country_code=code,
                        available=available,
                        direct_co2_g_per_kwh=co2_value,
                        radioactive_waste_mg_per_kwh=waste_value,
                        shares=shares,
                        untracked_share=untracked,
                        source_row=row_number,
                        residual_mix_row=row_number,
                        source_sheet=residual.title,
                    )
                )
            if not rows:
                raise PermanentIngestionError("AIB workbook contains no country rows")
            self.metrics = {
                "countries": len(rows),
                "available_residual_mix": sum(row.available for row in rows),
                "full_disclosure_no_factor": sum(not row.available for row in rows),
            }
            return tuple(rows)
        finally:
            workbook.close()

    @classmethod
    def _assert_residual_headers(cls, sheet: Any) -> None:
        values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        actual = tuple(cls._header_key(value) for value in values[: len(RESIDUAL_MIX_HEADERS)])
        expected = tuple(cls._header_key(value) for value in RESIDUAL_MIX_HEADERS)
        if actual[0] not in {"", "countrycode"} or actual[1:] != expected[1:]:
            raise PermanentIngestionError(
                f"AIB {sheet.title} schema changed; parser version must be updated"
            )

    @staticmethod
    def _header_key(value: object) -> str:
        if value is None:
            return ""
        return re.sub(r"[^a-z0-9]+", "", str(value).casefold())

    @classmethod
    def _assert_headers(cls, sheet: Any, expected: tuple[str | None, ...]) -> None:
        values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        actual = tuple(cls._header(value) for value in values[: len(expected)])
        if actual != expected:
            raise PermanentIngestionError(
                f"AIB {sheet.title} schema changed; parser version must be updated"
            )

    @classmethod
    def _residual_rows(
        cls, sheet: Any
    ) -> dict[str, tuple[int, dict[str, Decimal], Decimal | None]]:
        rows: dict[str, tuple[int, dict[str, Decimal], Decimal | None]] = {}
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            code = cls._code(values[0] if values else None)
            if code is None:
                continue
            if code in rows:
                raise PermanentIngestionError(f"duplicate AIB country code: {code}")
            shares: dict[str, Decimal] = {}
            for name, value in zip(SHARE_NAMES, values[1:15], strict=True):
                parsed = cls._share(value, row_number)
                if parsed is not None:
                    shares[name] = parsed
            rows[code] = (row_number, shares, cls._share(values[15], row_number))
        return rows

    @classmethod
    def _attribute_rows(cls, sheet: Any, label: str) -> dict[str, tuple[int, Decimal | None]]:
        rows: dict[str, tuple[int, Decimal | None]] = {}
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            code = cls._code(values[0] if values else None)
            if code is None:
                continue
            if code in rows:
                raise PermanentIngestionError(f"duplicate AIB {label} country code: {code}")
            rows[code] = (row_number, cls._decimal(values[2], row_number))
        return rows

    @staticmethod
    def _header(value: object) -> str | None:
        if value is None:
            return None
        return " ".join(str(value).split())

    @staticmethod
    def _code(value: object) -> str | None:
        if value is None:
            return None
        code = str(value).strip().upper()
        return code if len(code) == 2 and code.isalpha() else None

    @classmethod
    def _share(cls, value: object, row_number: int) -> Decimal | None:
        if isinstance(value, str) and value.strip().endswith("%"):
            parsed = cls._decimal(value.strip()[:-1], row_number)
            return parsed / Decimal(100) if parsed is not None else None
        return cls._decimal(value, row_number)

    @staticmethod
    def _decimal(value: object, row_number: int) -> Decimal | None:
        if value is None or value == "NA" or isinstance(value, time):
            return None
        try:
            return Decimal(str(value).strip().replace(",", "."))
        except InvalidOperation as error:
            raise PermanentIngestionError(
                f"invalid AIB numeric value at row {row_number}: {value}"
            ) from error
