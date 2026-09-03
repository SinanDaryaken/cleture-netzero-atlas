from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from itertools import pairwise
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, Field

from atlas.domain.enums import FactorIntendedUse, FactorValueKind
from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError

NumericCell = Decimal | str | None
ObservationKind = Literal["conversion_factor", "technical_property"]


class GhgProtocolFactorRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    source_sheet: str
    source_table: str
    source_row: int = Field(ge=1)
    source_key: str
    name: str
    category: str
    activity_type: str
    geography_code: str
    geography_name: str
    denominator_unit: str
    co2: NumericCell = None
    co2_unit: str | None = None
    ch4: NumericCell = None
    ch4_unit: str | None = None
    n2o: NumericCell = None
    n2o_unit: str | None = None
    direct_co2e: NumericCell = None
    direct_co2e_unit: str | None = None
    factor_value_kind: FactorValueKind
    intended_use: FactorIntendedUse = FactorIntendedUse.INVENTORY
    reference_year: int | None = None
    scope: str
    lifecycle_stage: str
    system_boundary: str
    original_source: str
    attributes: dict[str, Any] = Field(default_factory=dict)

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data={"record_type": "factor", **self.model_dump(mode="json")},
            source_sheet=self.source_sheet,
            source_table=self.source_table,
            source_row=self.source_row,
        )


class GhgProtocolObservationRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    observation_kind: ObservationKind
    source_sheet: str
    source_table: str
    source_row: int = Field(ge=1)
    source_column: int = Field(ge=1)
    name: str
    value: Decimal
    unit: str
    reference_year: int | None = None
    category: str
    attributes: dict[str, Any] = Field(default_factory=dict)

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data={"record_type": self.observation_kind, **self.model_dump(mode="json")},
            source_sheet=self.source_sheet,
            source_table=self.source_table,
            source_row=self.source_row,
        )


class GhgProtocolDocument(BaseModel):
    model_config = ConfigDict(frozen=True)

    version: str
    factors: tuple[GhgProtocolFactorRow, ...]
    observations: tuple[GhgProtocolObservationRow, ...]

    @property
    def parsed_records(self) -> tuple[ParsedRecord, ...]:
        return tuple(
            [row.as_parsed_record() for row in self.factors]
            + [row.as_parsed_record() for row in self.observations]
        )


class GhgProtocolParser:
    """Parse the audited GHG Protocol Cross-sector Emission Factors v2.0 workbook."""

    VERSION = "2.0"
    EXPECTED_SHEETS = (
        "Welcome!",
        "Table of Contents",
        "Stationary Combustion",
        "Mobile Combustion - Fuel Use",
        "Mobile Combustion - Distance",
        "Electricity US",
        "Electricity CN, TW, BR, TH, UK",
        "Mobile Combustion - Freight",
        "Mobile Combustion - Public",
        "Abbreviations and Conversions",
        "Revision History",
    )
    _US_ELECTRICITY_TABLE_ROWS = (3, 35, 67, 99, 131, 162, 193, 224, 255, 286, 317, 348)

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(self, path: Path) -> GhgProtocolDocument:
        try:
            workbook = load_workbook(path, read_only=False, data_only=True)
            self._assert_contract(workbook)
            factors = [
                *self._stationary(workbook["Stationary Combustion"]),
                *self._mobile_fuel(workbook["Mobile Combustion - Fuel Use"]),
                *self._mobile_distance(workbook["Mobile Combustion - Distance"]),
                *self._electricity_us(workbook["Electricity US"]),
                *self._electricity_other(workbook["Electricity CN, TW, BR, TH, UK"]),
                *self._freight(workbook["Mobile Combustion - Freight"]),
                *self._public_transport(workbook["Mobile Combustion - Public"]),
            ]
            observations = [
                *self._fuel_economy(workbook["Mobile Combustion - Distance"]),
                *self._conversions(workbook["Abbreviations and Conversions"]),
            ]
        except PermanentIngestionError:
            raise
        except Exception as error:
            raise PermanentIngestionError(
                "GHG Protocol asset is not a readable Cross-sector v2.0 workbook"
            ) from error

        if not factors or not observations:
            raise PermanentIngestionError("GHG Protocol workbook produced no records")
        self.metrics = {
            "parsed_rows": len(factors) + len(observations),
            "parsed_factor_rows": len(factors),
            "parsed_observations": len(observations),
            "range_valued_factor_rows": sum(self._has_range(row) for row in factors),
        }
        return GhgProtocolDocument(
            version=self.VERSION,
            factors=tuple(factors),
            observations=tuple(observations),
        )

    @classmethod
    def _assert_contract(cls, workbook: Any) -> None:
        if tuple(workbook.sheetnames) != cls.EXPECTED_SHEETS:
            raise PermanentIngestionError("GHG Protocol Cross-sector workbook sheets changed")
        welcome = str(workbook["Welcome!"]["B10"].value or "")
        revision = workbook["Revision History"]
        if "March 2024" not in welcome or str(revision["B6"].value or "") != "2.0":
            raise PermanentIngestionError("GHG Protocol Cross-sector workbook version changed")

    def _stationary(self, sheet: Worksheet) -> list[GhgProtocolFactorRow]:
        categories: dict[int, str] = {}
        current = ""
        for row in range(6, 60):
            current = self._text(sheet.cell(row, 2).value) or current
            categories[row] = current
        rows: list[GhgProtocolFactorRow] = []
        bases = ((5, "TJ"), (6, "tonne"), (9, "litre"), (10, "m3"))
        for row in range(6, 60):
            fuel = self._text(sheet.cell(row, 3).value)
            if not fuel:
                continue
            ch4_row = row + 65
            n2o_row = row + 130
            for column, denominator in bases:
                co2 = self._cell(sheet.cell(row, column).value)
                ch4 = self._cell(sheet.cell(ch4_row, column).value)
                n2o = self._cell(sheet.cell(n2o_row, column).value)
                if co2 is None and ch4 is None and n2o is None:
                    continue
                rows.append(
                    self._factor(
                        sheet,
                        table="Tables 1-3 - Stationary combustion",
                        row=row,
                        source_key=f"stationary-{fuel}-{denominator}",
                        name=f"{fuel} stationary combustion",
                        category=categories[row],
                        activity_type="stationary-combustion",
                        geography_code="GLOBAL",
                        geography_name="Global",
                        denominator_unit=denominator,
                        co2=co2,
                        co2_unit=f"kg/{denominator}",
                        ch4=ch4,
                        ch4_unit=f"kg/{denominator}",
                        n2o=n2o,
                        n2o_unit=f"kg/{denominator}",
                        kind=FactorValueKind.CO2E_TOTAL,
                        scope="scope 1",
                        lifecycle="tank-to-wheel",
                        boundary="direct fuel combustion",
                        original_source=(
                            "2006 IPCC Guidelines for National Greenhouse Gas Inventories"
                        ),
                    )
                )
        return rows

    def _mobile_fuel(self, sheet: Worksheet) -> list[GhgProtocolFactorRow]:
        rows: list[GhgProtocolFactorRow] = []
        for row in range(4, 37):
            region = self._text(sheet.cell(row, 2).value)
            fuel = self._text(sheet.cell(row, 3).value)
            unit = self._text(sheet.cell(row, 6).value)
            if not region or not fuel or not unit:
                continue
            denominator = unit.split("/", 1)[1]
            for column, label, intended in (
                (4, "fossil", FactorIntendedUse.INVENTORY),
                (5, "biogenic", FactorIntendedUse.CALCULATION_INPUT),
            ):
                value = self._cell(sheet.cell(row, column).value)
                if value is None:
                    continue
                code, name = self._region(region)
                rows.append(
                    self._factor(
                        sheet,
                        table="Table 1 - Mobile fuel CO2",
                        row=row,
                        source_key=f"mobile-fuel-{region}-{fuel}-{label}",
                        name=f"{fuel} mobile combustion ({label} CO2)",
                        category="Mobile combustion - fuel use",
                        activity_type="mobile-combustion-fuel",
                        geography_code=code,
                        geography_name=name,
                        denominator_unit=denominator,
                        co2=value,
                        co2_unit=unit,
                        kind=FactorValueKind.CO2_ONLY,
                        intended=intended,
                        scope="outside scopes" if label == "biogenic" else "scope 1",
                        lifecycle="tank-to-wheel",
                        boundary="direct mobile fuel combustion",
                        original_source="IPCC, US EPA, or UK Government GHG Conversion Factors",
                        attributes={"carbon_origin": label},
                    )
                )
        for row in range(46, 99):
            region = self._text(sheet.cell(row, 2).value)
            fuel = self._text(sheet.cell(row, 3).value)
            vehicle = self._text(sheet.cell(row, 4).value)
            if not region or not fuel or not vehicle:
                continue
            code, name = self._region(region)
            rows.append(
                self._factor(
                    sheet,
                    table="Table 2 - Mobile fuel CH4 and N2O",
                    row=row,
                    source_key=(
                        f"mobile-fuel-nonco2-{region}-{fuel}-{vehicle}-{sheet.cell(row, 5).value}"
                    ),
                    name=f"{fuel} - {vehicle} mobile combustion non-CO2",
                    category="Mobile combustion - fuel use",
                    activity_type="mobile-combustion-fuel",
                    geography_code=code,
                    geography_name=name,
                    denominator_unit=self._denominator(sheet.cell(row, 7).value),
                    ch4=self._cell(sheet.cell(row, 6).value),
                    ch4_unit=self._text(sheet.cell(row, 7).value),
                    n2o=self._cell(sheet.cell(row, 8).value),
                    n2o_unit=self._text(sheet.cell(row, 9).value),
                    kind=FactorValueKind.NON_CO2_CO2E,
                    scope="scope 1",
                    lifecycle="tank-to-wheel",
                    boundary="direct mobile fuel combustion",
                    original_source="2006 IPCC Guidelines or US EPA Emission Factors Hub",
                    attributes={"engine_type": self._text(sheet.cell(row, 5).value)},
                )
            )
        return rows

    def _mobile_distance(self, sheet: Worksheet) -> list[GhgProtocolFactorRow]:
        rows: list[GhgProtocolFactorRow] = []
        for row in range(18, 178):
            region = self._text(sheet.cell(row, 2).value)
            vehicle = self._text(sheet.cell(row, 3).value)
            fuel = self._text(sheet.cell(row, 5).value)
            if not region or not vehicle or not fuel:
                continue
            code, name = self._region(region)
            rows.append(
                self._factor(
                    sheet,
                    table="Table 1 - Mobile distance CH4 and N2O",
                    row=row,
                    source_key=(
                        f"mobile-distance-nonco2-{region}-{vehicle}-"
                        f"{sheet.cell(row, 4).value}-{fuel}"
                    ),
                    name=f"{vehicle} - {fuel} distance non-CO2",
                    category="Mobile combustion - vehicle distance",
                    activity_type="vehicle-distance",
                    geography_code=code,
                    geography_name=name,
                    denominator_unit=self._denominator(sheet.cell(row, 7).value),
                    ch4=self._cell(sheet.cell(row, 6).value),
                    ch4_unit=self._text(sheet.cell(row, 7).value),
                    n2o=self._cell(sheet.cell(row, 8).value),
                    n2o_unit=self._text(sheet.cell(row, 9).value),
                    kind=FactorValueKind.NON_CO2_CO2E,
                    scope="scope 1",
                    lifecycle="tank-to-wheel",
                    boundary="direct vehicle operation",
                    original_source="2006 IPCC Guidelines or US EPA Emission Factors Hub",
                    attributes={"vehicle_year": self._text(sheet.cell(row, 4).value)},
                )
            )
        for row in range(183, 236):
            region = self._text(sheet.cell(row, 2).value)
            vehicle = self._text(sheet.cell(row, 3).value)
            if not region or not vehicle:
                continue
            code, name = self._region(region)
            rows.append(
                self._factor(
                    sheet,
                    table="Table 2 - UK mobile distance",
                    row=row,
                    source_key=(
                        f"mobile-distance-uk-{vehicle}-{sheet.cell(row, 4).value}-"
                        f"{sheet.cell(row, 5).value}-{sheet.cell(row, 6).value}"
                    ),
                    name=f"{vehicle} - {self._text(sheet.cell(row, 4).value) or 'average'}",
                    category="Mobile combustion - vehicle distance",
                    activity_type="vehicle-distance",
                    geography_code=code,
                    geography_name=name,
                    denominator_unit="km",
                    co2=self._cell(sheet.cell(row, 7).value),
                    co2_unit="kg/km",
                    ch4=self._cell(sheet.cell(row, 8).value),
                    ch4_unit="g/km",
                    n2o=self._cell(sheet.cell(row, 9).value),
                    n2o_unit="g/km",
                    kind=FactorValueKind.CO2E_TOTAL,
                    scope="scope 1",
                    lifecycle="tank-to-wheel",
                    boundary="direct vehicle operation",
                    original_source=(
                        "UK Government GHG Conversion Factors for Company Reporting 2023"
                    ),
                    attributes={
                        "size": self._text(sheet.cell(row, 4).value),
                        "weight_laden": self._text(sheet.cell(row, 5).value),
                        "fuel": self._text(sheet.cell(row, 6).value),
                    },
                )
            )
        return rows

    def _electricity_us(self, sheet: Worksheet) -> list[GhgProtocolFactorRow]:
        rows: list[GhgProtocolFactorRow] = []
        anchors = (*self._US_ELECTRICITY_TABLE_ROWS, sheet.max_row + 1)
        for table_number, (start, stop) in enumerate(pairwise(anchors), 1):
            title = self._text(sheet.cell(start, 2).value)
            match = re.search(r"Year (\d{4})", title)
            if not match:
                raise PermanentIngestionError(
                    f"GHG Protocol eGRID table title changed at row {start}"
                )
            year = int(match.group(1))
            for row in range(start + 3, stop):
                region = self._text(sheet.cell(row, 2).value)
                co2 = self._cell(sheet.cell(row, 3).value)
                ch4 = self._cell(sheet.cell(row, 4).value)
                n2o = self._cell(sheet.cell(row, 5).value)
                if not region or (co2 is None and ch4 is None and n2o is None):
                    continue
                rows.append(
                    self._factor(
                        sheet,
                        table=f"Table {table_number} - eGRID {year}",
                        row=row,
                        source_key=f"electricity-us-{year}-{region}",
                        name=f"{region} grid electricity {year}",
                        category="Purchased electricity",
                        activity_type="purchased-electricity",
                        geography_code=region,
                        geography_name=region,
                        denominator_unit="MWh",
                        co2=co2,
                        co2_unit="lb/MWh",
                        ch4=ch4,
                        ch4_unit="lb/GWh",
                        n2o=n2o,
                        n2o_unit="lb/GWh",
                        kind=FactorValueKind.CO2E_TOTAL,
                        reference_year=year,
                        scope="scope 2",
                        lifecycle="generation",
                        boundary="grid average electricity generation",
                        original_source="US EPA eGRID",
                    )
                )
        return rows

    def _electricity_other(self, sheet: Worksheet) -> list[GhgProtocolFactorRow]:
        specs = (
            (1, range(6, 7), "CN", "China", "t/MWh", FactorValueKind.CO2_ONLY),
            (2, range(11, 29), "TW", "Taiwan", "kg/kWh", FactorValueKind.CO2E_TOTAL),
            (3, range(36, 44), "BR", "Brazil", "t/MWh", FactorValueKind.CO2_ONLY),
            (4, range(49, 61), "TH", "Thailand", "kg/kWh", FactorValueKind.CO2_ONLY),
        )
        rows: list[GhgProtocolFactorRow] = []
        for table, source_rows, code, country, unit, kind in specs:
            for row in source_rows:
                year = self._year(sheet.cell(row, 2).value)
                value = self._cell(sheet.cell(row, 3).value)
                if year is None or value is None:
                    continue
                kwargs: dict[str, Any]
                if kind == FactorValueKind.CO2E_TOTAL:
                    kwargs = {"direct_co2e": value, "direct_co2e_unit": unit}
                else:
                    kwargs = {"co2": value, "co2_unit": unit}
                rows.append(
                    self._factor(
                        sheet,
                        table=f"Table {table} - {country} electricity",
                        row=row,
                        source_key=f"electricity-{code}-{year}",
                        name=f"{country} grid electricity {year}",
                        category="Purchased electricity",
                        activity_type="purchased-electricity",
                        geography_code=code,
                        geography_name=country,
                        denominator_unit=unit.split("/", 1)[1],
                        kind=kind,
                        reference_year=year,
                        scope="scope 2",
                        lifecycle="generation",
                        boundary="national grid average electricity generation",
                        original_source=f"{country} national electricity factor publisher",
                        **kwargs,
                    )
                )
        row = 66
        rows.append(
            self._factor(
                sheet,
                table="Table 5 - United Kingdom electricity",
                row=row,
                source_key="electricity-GB-2023",
                name="United Kingdom grid electricity 2023",
                category="Purchased electricity",
                activity_type="purchased-electricity",
                geography_code="GB",
                geography_name="United Kingdom",
                denominator_unit="kWh",
                co2=self._cell(sheet.cell(row, 3).value),
                co2_unit="kg/kWh",
                ch4=self._cell(sheet.cell(row, 4).value),
                ch4_unit="kg/kWh",
                n2o=self._cell(sheet.cell(row, 5).value),
                n2o_unit="kg/kWh",
                kind=FactorValueKind.CO2E_TOTAL,
                reference_year=2023,
                scope="scope 2",
                lifecycle="generation",
                boundary="national grid average electricity generation",
                original_source="UK Government GHG Conversion Factors for Company Reporting 2023",
                attributes={"gas_components_reverse_calculated_from_ar5_total": True},
            )
        )
        return rows

    def _freight(self, sheet: Worksheet) -> list[GhgProtocolFactorRow]:
        co2_rows = [row for row in range(5, 64) if self._cell(sheet.cell(row, 7).value) is not None]
        gas_rows = [
            row for row in range(79, 138) if self._cell(sheet.cell(row, 7).value) is not None
        ]
        if len(co2_rows) != len(gas_rows):
            raise PermanentIngestionError("GHG Protocol freight tables no longer align")
        rows: list[GhgProtocolFactorRow] = []
        for co2_row, gas_row in zip(co2_rows, gas_rows, strict=True):
            region = self._text(sheet.cell(co2_row, 2).value)
            vehicle = self._text(sheet.cell(co2_row, 3).value)
            code, name = self._region(region)
            unit = self._text(sheet.cell(co2_row, 8).value)
            rows.append(
                self._factor(
                    sheet,
                    table="Tables 1-2 - Freight transport",
                    row=co2_row,
                    source_key=(
                        f"freight-{region}-{vehicle}-{sheet.cell(co2_row, 4).value}-"
                        f"{sheet.cell(co2_row, 5).value}-{sheet.cell(co2_row, 6).value}"
                    ),
                    name=" - ".join(
                        filter(None, [vehicle, self._text(sheet.cell(co2_row, 4).value)])
                    ),
                    category="Freight transport",
                    activity_type="freight-transport",
                    geography_code=code,
                    geography_name=name,
                    denominator_unit=self._denominator(unit),
                    co2=self._cell(sheet.cell(co2_row, 7).value),
                    co2_unit=unit,
                    ch4=self._cell(sheet.cell(gas_row, 7).value),
                    ch4_unit=self._text(sheet.cell(gas_row, 8).value),
                    n2o=self._cell(sheet.cell(gas_row, 9).value),
                    n2o_unit=self._text(sheet.cell(gas_row, 10).value),
                    kind=FactorValueKind.CO2E_TOTAL,
                    scope="scope 3",
                    lifecycle="tank-to-wheel",
                    boundary="third-party freight transport operation",
                    original_source=(
                        "UK Government GHG Conversion Factors or US EPA Emission Factors Hub"
                    ),
                    attributes={
                        "weight_class": self._text(sheet.cell(co2_row, 5).value),
                        "fuel": self._text(sheet.cell(co2_row, 6).value),
                        "component_row": gas_row,
                    },
                )
            )
        return rows

    def _public_transport(self, sheet: Worksheet) -> list[GhgProtocolFactorRow]:
        rows: list[GhgProtocolFactorRow] = []
        basis = "public-passenger-distance"
        for row in range(7, 41):
            marker = self._text(sheet.cell(row, 2).value)
            if marker in {"Passenger-Distance", "Vehicle-Distance"}:
                basis = f"public-{marker.lower()}"
                continue
            region = marker
            transport = self._text(sheet.cell(row, 3).value)
            if not region or not transport:
                continue
            code, name = self._region(region)
            unit = self._text(sheet.cell(row, 6).value)
            rows.append(
                self._factor(
                    sheet,
                    table="Table 1 - Public transport",
                    row=row,
                    source_key=f"public-{region}-{transport}-{sheet.cell(row, 4).value}",
                    name=" - ".join(
                        filter(None, [transport, self._text(sheet.cell(row, 4).value)])
                    ),
                    category="Public transport",
                    activity_type=basis,
                    geography_code=code,
                    geography_name=name,
                    denominator_unit=self._denominator(unit),
                    co2=self._cell(sheet.cell(row, 5).value),
                    co2_unit=unit,
                    ch4=self._cell(sheet.cell(row, 7).value),
                    ch4_unit=self._text(sheet.cell(row, 8).value),
                    n2o=self._cell(sheet.cell(row, 9).value),
                    n2o_unit=self._text(sheet.cell(row, 10).value),
                    kind=FactorValueKind.CO2E_TOTAL,
                    scope="scope 3",
                    lifecycle="tank-to-wheel",
                    boundary="third-party passenger transport operation",
                    original_source=(
                        "UK Government GHG Conversion Factors or US EPA Emission Factors Hub"
                    ),
                )
            )
        return rows

    def _fuel_economy(self, sheet: Worksheet) -> list[GhgProtocolObservationRow]:
        rows: list[GhgProtocolObservationRow] = []
        for row in range(6, 12):
            vehicle = self._text(sheet.cell(row, 2).value)
            for column, unit in ((4, "mile/US gallon"), (5, "km/l")):
                value = self._decimal(sheet.cell(row, column).value)
                if vehicle and value is not None:
                    rows.append(
                        GhgProtocolObservationRow(
                            observation_kind="technical_property",
                            source_sheet=sheet.title,
                            source_table="Reference Table - Average Fuel Economy",
                            source_row=row,
                            source_column=column,
                            name=f"{vehicle} average fuel economy",
                            value=value,
                            unit=unit,
                            reference_year=2021,
                            category="Vehicle fuel economy",
                            attributes={
                                "source_publication": "FHWA Highway Statistics 2021 (January 2024)"
                            },
                        )
                    )
        return rows

    def _conversions(self, sheet: Worksheet) -> list[GhgProtocolObservationRow]:
        rows: list[GhgProtocolObservationRow] = []
        for table, header_row, start_row, end_row, end_column in (
            ("Energy conversions", 5, 6, 11, 8),
            ("Volume conversions", 15, 16, 21, 8),
            ("Weight/mass conversions", 25, 26, 31, 8),
            ("Length/distance conversions", 35, 36, 43, 10),
        ):
            for row in range(start_row, end_row + 1):
                source_unit = self._text(sheet.cell(row, 2).value)
                for column in range(3, end_column + 1):
                    target_unit = self._text(sheet.cell(header_row, column).value)
                    value = self._decimal(sheet.cell(row, column).value)
                    if source_unit and target_unit and value is not None:
                        rows.append(
                            GhgProtocolObservationRow(
                                observation_kind="conversion_factor",
                                source_sheet=sheet.title,
                                source_table=table,
                                source_row=row,
                                source_column=column,
                                name=f"{source_unit} to {target_unit}",
                                value=value,
                                unit=f"{target_unit}/{source_unit}",
                                category=table,
                                attributes={"source_unit": source_unit, "target_unit": target_unit},
                            )
                        )
        return rows

    @staticmethod
    def _factor(
        sheet: Worksheet,
        *,
        table: str,
        row: int,
        source_key: str,
        name: str,
        category: str,
        activity_type: str,
        geography_code: str,
        geography_name: str,
        denominator_unit: str,
        kind: FactorValueKind,
        scope: str,
        lifecycle: str,
        boundary: str,
        original_source: str,
        co2: NumericCell = None,
        co2_unit: str | None = None,
        ch4: NumericCell = None,
        ch4_unit: str | None = None,
        n2o: NumericCell = None,
        n2o_unit: str | None = None,
        direct_co2e: NumericCell = None,
        direct_co2e_unit: str | None = None,
        intended: FactorIntendedUse = FactorIntendedUse.INVENTORY,
        reference_year: int | None = None,
        attributes: dict[str, Any] | None = None,
    ) -> GhgProtocolFactorRow:
        return GhgProtocolFactorRow(
            source_sheet=sheet.title,
            source_table=table,
            source_row=row,
            source_key=source_key,
            name=name.strip(),
            category=category,
            activity_type=activity_type,
            geography_code=geography_code,
            geography_name=geography_name,
            denominator_unit=denominator_unit,
            co2=co2,
            co2_unit=co2_unit,
            ch4=ch4,
            ch4_unit=ch4_unit,
            n2o=n2o,
            n2o_unit=n2o_unit,
            direct_co2e=direct_co2e,
            direct_co2e_unit=direct_co2e_unit,
            factor_value_kind=kind,
            intended_use=intended,
            reference_year=reference_year,
            scope=scope,
            lifecycle_stage=lifecycle,
            system_boundary=boundary,
            original_source=original_source,
            attributes=attributes or {},
        )

    @staticmethod
    def _region(value: str) -> tuple[str, str]:
        cleaned = re.sub(r"\d+$", "", value).strip()
        return {
            "Other": ("GLOBAL", "Global"),
            "US": ("US", "United States"),
            "UK": ("GB", "United Kingdom"),
        }.get(cleaned, (cleaned, cleaned))

    @staticmethod
    def _denominator(value: Any) -> str:
        text = str(value or "").strip()
        return text.split("/", 1)[1] if "/" in text else text

    @staticmethod
    def _text(value: Any) -> str:
        return " ".join(str(value or "").split()).strip()

    @classmethod
    def _cell(cls, value: Any) -> NumericCell:
        decimal = cls._decimal(value)
        if decimal is not None:
            return decimal
        text = cls._text(value)
        return text or None

    @staticmethod
    def _decimal(value: Any) -> Decimal | None:
        if value is None or isinstance(value, bool):
            return None
        try:
            return Decimal(str(value).replace(",", ""))
        except (InvalidOperation, ValueError):
            return None

    @classmethod
    def _year(cls, value: Any) -> int | None:
        decimal = cls._decimal(value)
        return int(decimal) if decimal is not None else None

    @staticmethod
    def _has_range(row: GhgProtocolFactorRow) -> bool:
        return any(isinstance(value, str) for value in (row.co2, row.ch4, row.n2o, row.direct_co2e))
