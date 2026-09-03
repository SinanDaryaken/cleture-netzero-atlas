from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import ClassVar, Literal

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, Field

from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError

FactorGeographyKind = Literal["country", "regional_average"]
ParameterType = Literal[
    "exchange_rate",
    "purchaser_producer_conversion",
    "sector_price_index",
]


class OpenCedaFactorRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    release_year: int
    base_year: int
    geography_code: str = Field(min_length=2)
    geography_name: str = Field(min_length=1)
    geography_kind: FactorGeographyKind
    applicable_country_codes: tuple[str, ...] = ()
    sector_code: str = Field(min_length=1)
    sector_name: str = Field(min_length=1)
    value: Decimal
    unit: str = Field(min_length=1)
    price_type: str = Field(min_length=1)
    currency: str = Field(min_length=1)
    source_sheet: str = Field(min_length=1)
    source_row: int = Field(ge=1)
    source_column: int = Field(ge=1)

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data={"record_type": "factor", **self.model_dump(mode="json")},
            source_sheet=self.source_sheet,
            source_row=self.source_row,
            source_table=self.geography_kind,
        )


class OpenCedaParameterRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    release_year: int
    base_year: int
    parameter_type: ParameterType
    name: str = Field(min_length=1)
    value: Decimal
    unit: str = Field(min_length=1)
    reference_year: int | None = None
    country_code: str | None = None
    country_name: str | None = None
    sector_code: str | None = None
    sector_name: str | None = None
    source_sheet: str = Field(min_length=1)
    source_row: int = Field(ge=1)
    source_column: int = Field(ge=1)

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data={"record_type": "calculation_parameter", **self.model_dump(mode="json")},
            source_sheet=self.source_sheet,
            source_row=self.source_row,
            source_table=self.parameter_type,
        )


@dataclass(frozen=True)
class OpenCedaWorkbook:
    release_year: int
    base_year: int
    published_at: datetime
    factors: tuple[OpenCedaFactorRow, ...]
    parameters: tuple[OpenCedaParameterRow, ...]


class OpenCedaParser:
    """Parse source matrices without treating adjustment tables as emission factors."""

    _regional_aliases: ClassVar[dict[str, str]] = {
        "Australian and New Zealand": "Australia and New Zealand",
        "Carribean": "Caribbean",
        "South Eastern Asia": "South-Eastern Asia",
    }

    def __init__(self, *, expected_sector_count: int = 400, expected_country_count: int = 149):
        self.expected_sector_count = expected_sector_count
        self.expected_country_count = expected_country_count
        self.metrics: dict[str, int] = {}

    def parse(self, path: Path, *, release_year: int) -> OpenCedaWorkbook:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as error:
            raise PermanentIngestionError(
                "Open CEDA asset is not a readable XLSX workbook"
            ) from error

        if release_year == 2024:
            result = self._parse_2024(workbook)
        elif release_year == 2025:
            result = self._parse_2025(workbook)
        else:
            raise PermanentIngestionError(f"Open CEDA schema is not registered for {release_year}")

        zero_values = sum(row.value == 0 for row in result.factors)
        country_matrix_factors = sum(row.geography_kind == "country" for row in result.factors)
        regional_factors = len(result.factors) - country_matrix_factors
        self.metrics = {
            "parsed_rows": len(result.factors) + len(result.parameters),
            "parsed_factor_cells": len(result.factors),
            "country_matrix_factors": country_matrix_factors,
            "regional_fallback_factors": regional_factors,
            "calculation_parameters": len(result.parameters),
            "zero_value_factors": zero_values,
            "reference_year": result.base_year,
            "release_year": result.release_year,
        }
        return result

    def _parse_2024(self, workbook: object) -> OpenCedaWorkbook:
        self._require_sheets(
            workbook,
            {
                "Open CEDA",
                "Exchange rates",
                "Purchaser - producer conversion",
                "Sector level Price Index",
                "Metadata",
            },
        )
        factor_sheet = workbook["Open CEDA"]  # type: ignore[index]
        version = str(factor_sheet.cell(13, 4).value or "")
        license_name = str(factor_sheet.cell(14, 4).value or "")
        if version != "CEDA 2024" or license_name != "CC BY-SA 4.0":
            raise PermanentIngestionError("Open CEDA 2024 release metadata changed")
        base_year = self._year(factor_sheet.cell(24, 3).value, "base year")
        price_type = str(factor_sheet.cell(25, 3).value or "")
        currency = str(factor_sheet.cell(26, 3).value or "")
        factors = self._parse_factor_matrix(
            factor_sheet,
            release_year=2024,
            base_year=base_year,
            price_type=price_type,
            currency=currency,
            name_row=27,
            code_row=28,
            data_start_row=29,
            first_factor_column=5,
            geography_code_column=2,
            geography_name_column=3,
            unit_column=4,
            geography_kind="country",
        )
        parameters = self._parse_parameters(workbook, release_year=2024, base_year=base_year)
        return OpenCedaWorkbook(
            release_year=2024,
            base_year=base_year,
            published_at=self._published_at(factor_sheet.cell(12, 4).value),
            factors=factors,
            parameters=parameters,
        )

    def _parse_2025(self, workbook: object) -> OpenCedaWorkbook:
        self._require_sheets(
            workbook,
            {
                "Cover",
                "GHG_t_Raw",
                "Regional Average EFs",
                "Exchange rates",
                "Purchaser - producer conversion",
                "Sector level Price Index",
                "Country to region mapping",
                "Metadata",
            },
        )
        cover = workbook["Cover"]  # type: ignore[index]
        factor_sheet = workbook["GHG_t_Raw"]  # type: ignore[index]
        version = str(cover.cell(10, 4).value or "")
        license_name = str(cover.cell(11, 4).value or "")
        if version != "CEDA 2025" or license_name != "CC BY-SA 4.0":
            raise PermanentIngestionError("Open CEDA 2025 release metadata changed")
        base_year = self._year(factor_sheet.cell(1, 2).value, "base year")
        price_type = str(factor_sheet.cell(2, 2).value or "")
        currency = str(factor_sheet.cell(2, 4).value or "")
        factors = list(
            self._parse_factor_matrix(
                factor_sheet,
                release_year=2025,
                base_year=base_year,
                price_type=price_type,
                currency=currency,
                name_row=3,
                code_row=4,
                data_start_row=5,
                first_factor_column=4,
                geography_code_column=1,
                geography_name_column=2,
                unit_column=3,
                geography_kind="country",
            )
        )
        country_to_region = self._country_to_region(workbook["Country to region mapping"])  # type: ignore[index]
        factors.extend(
            self._parse_regional_matrix(
                workbook["Regional Average EFs"],  # type: ignore[index]
                release_year=2025,
                base_year=base_year,
                country_to_region=country_to_region,
            )
        )
        parameters = self._parse_parameters(workbook, release_year=2025, base_year=base_year)
        return OpenCedaWorkbook(
            release_year=2025,
            base_year=base_year,
            published_at=self._published_at(cover.cell(9, 4).value),
            factors=tuple(factors),
            parameters=parameters,
        )

    def _parse_factor_matrix(
        self,
        sheet: Worksheet,
        *,
        release_year: int,
        base_year: int,
        price_type: str,
        currency: str,
        name_row: int,
        code_row: int,
        data_start_row: int,
        first_factor_column: int,
        geography_code_column: int,
        geography_name_column: int,
        unit_column: int,
        geography_kind: FactorGeographyKind,
    ) -> tuple[OpenCedaFactorRow, ...]:
        sectors = self._sectors(sheet, name_row, code_row, first_factor_column)
        rows: list[OpenCedaFactorRow] = []
        geography_count = 0
        maximum_column = first_factor_column + self.expected_sector_count - 1
        for source_row, values in enumerate(
            sheet.iter_rows(
                min_row=data_start_row,
                max_col=maximum_column,
                values_only=True,
            ),
            start=data_start_row,
        ):
            geography_code = values[geography_code_column - 1]
            if geography_code is None:
                continue
            geography_count += 1
            geography_name = values[geography_name_column - 1]
            unit = str(values[unit_column - 1] or "")
            if unit != "kgCO2e/US Dollar":
                raise PermanentIngestionError(
                    f"Open CEDA factor unit changed at {sheet.title}!{source_row}: {unit!r}"
                )
            for source_column, sector_code, sector_name in sectors:
                value = values[source_column - 1]
                rows.append(
                    OpenCedaFactorRow(
                        release_year=release_year,
                        base_year=base_year,
                        geography_code=str(geography_code).strip(),
                        geography_name=str(geography_name or geography_code).strip(),
                        geography_kind=geography_kind,
                        sector_code=sector_code,
                        sector_name=sector_name,
                        value=self._decimal(value, sheet.title, source_row, source_column),
                        unit=unit,
                        price_type=price_type,
                        currency=currency,
                        source_sheet=sheet.title,
                        source_row=source_row,
                        source_column=source_column,
                    )
                )
        if geography_count != self.expected_country_count:
            raise PermanentIngestionError(
                f"Open CEDA country row count changed: {geography_count}; "
                f"expected {self.expected_country_count}"
            )
        return tuple(rows)

    def _parse_regional_matrix(
        self,
        sheet: Worksheet,
        *,
        release_year: int,
        base_year: int,
        country_to_region: dict[str, tuple[str, ...]],
    ) -> tuple[OpenCedaFactorRow, ...]:
        sectors = self._sectors(sheet, 3, 4, 3)
        rows: list[OpenCedaFactorRow] = []
        maximum_column = 2 + self.expected_sector_count
        for source_row, values in enumerate(
            sheet.iter_rows(min_row=5, max_col=maximum_column, values_only=True),
            start=5,
        ):
            source_name = values[0]
            if source_name is None:
                continue
            region_name = str(source_name).strip()
            mapping_name = self._regional_aliases.get(region_name, region_name)
            applicable = country_to_region.get(mapping_name)
            if not applicable:
                raise PermanentIngestionError(
                    f"Open CEDA regional geography has no country mapping: {region_name}"
                )
            unit = str(values[1] or "")
            if unit != "kgCO2e/US Dollar":
                raise PermanentIngestionError(
                    f"Open CEDA regional unit changed at row {source_row}: {unit!r}"
                )
            for source_column, sector_code, sector_name in sectors:
                rows.append(
                    OpenCedaFactorRow(
                        release_year=release_year,
                        base_year=base_year,
                        geography_code=self._slug(mapping_name),
                        geography_name=mapping_name,
                        geography_kind="regional_average",
                        applicable_country_codes=applicable,
                        sector_code=sector_code,
                        sector_name=sector_name,
                        value=self._decimal(
                            values[source_column - 1],
                            sheet.title,
                            source_row,
                            source_column,
                        ),
                        unit=unit,
                        price_type="Producer price",
                        currency="US Dollar",
                        source_sheet=sheet.title,
                        source_row=source_row,
                        source_column=source_column,
                    )
                )
        return tuple(rows)

    def _parse_parameters(
        self, workbook: object, *, release_year: int, base_year: int
    ) -> tuple[OpenCedaParameterRow, ...]:
        rows = [
            *self._parse_exchange_rates(
                workbook["Exchange rates"], release_year=release_year, base_year=base_year  # type: ignore[index]
            ),
            *self._parse_purchaser_conversion(
                workbook["Purchaser - producer conversion"],  # type: ignore[index]
                release_year=release_year,
                base_year=base_year,
            ),
            *self._parse_price_indices(
                workbook["Sector level Price Index"],  # type: ignore[index]
                release_year=release_year,
                base_year=base_year,
            ),
        ]
        return tuple(rows)

    def _parse_exchange_rates(
        self, sheet: Worksheet, *, release_year: int, base_year: int
    ) -> tuple[OpenCedaParameterRow, ...]:
        years: list[tuple[int, int]] = []
        header = next(sheet.iter_rows(min_row=4, max_row=4, max_col=32, values_only=True))
        column = 4
        while column <= len(header):
            value = header[column - 1]
            if value is None:
                break
            years.append((column, self._year(value, "exchange-rate year")))
            column += 1
        rows: list[OpenCedaParameterRow] = []
        maximum_column = years[-1][0]
        for source_row, values in enumerate(
            sheet.iter_rows(min_row=5, max_col=maximum_column, values_only=True),
            start=5,
        ):
            country_code = values[0]
            if country_code is None:
                continue
            country_name = str(values[1] or country_code)
            currency_name = str(values[2] or "local currency")
            for source_column, reference_year in years:
                value = values[source_column - 1]
                if value is None:
                    continue
                rows.append(
                    OpenCedaParameterRow(
                        release_year=release_year,
                        base_year=base_year,
                        parameter_type="exchange_rate",
                        name=f"{country_name} local currency units per USD",
                        value=self._decimal(value, sheet.title, source_row, source_column),
                        unit="LCU/USD",
                        reference_year=reference_year,
                        country_code=str(country_code),
                        country_name=country_name,
                        source_sheet=sheet.title,
                        source_row=source_row,
                        source_column=source_column,
                    )
                )
                rows[-1] = rows[-1].model_copy(
                    update={"name": f"{country_name} {currency_name} per USD"}
                )
        return tuple(rows)

    def _parse_purchaser_conversion(
        self, sheet: Worksheet, *, release_year: int, base_year: int
    ) -> tuple[OpenCedaParameterRow, ...]:
        sectors = self._sectors(sheet, 4, 5, 2)
        rows: list[OpenCedaParameterRow] = []
        values = next(
            sheet.iter_rows(
                min_row=6,
                max_row=6,
                max_col=1 + self.expected_sector_count,
                values_only=True,
            )
        )
        for source_column, sector_code, sector_name in sectors:
            rows.append(
                OpenCedaParameterRow(
                    release_year=release_year,
                    base_year=base_year,
                    parameter_type="purchaser_producer_conversion",
                    name=f"Purchaser/producer price conversion — {sector_name}",
                    value=self._decimal(
                        values[source_column - 1],
                        sheet.title,
                        6,
                        source_column,
                    ),
                    unit="ratio",
                    sector_code=sector_code,
                    sector_name=sector_name,
                    source_sheet=sheet.title,
                    source_row=6,
                    source_column=source_column,
                )
            )
        return tuple(rows)

    def _parse_price_indices(
        self, sheet: Worksheet, *, release_year: int, base_year: int
    ) -> tuple[OpenCedaParameterRow, ...]:
        sectors = self._sectors(sheet, 4, 5, 2)
        rows: list[OpenCedaParameterRow] = []
        for source_row, values in enumerate(
            sheet.iter_rows(
                min_row=6,
                max_col=1 + self.expected_sector_count,
                values_only=True,
            ),
            start=6,
        ):
            value = values[0]
            if value is None:
                continue
            try:
                reference_year = self._year(value, "price-index year")
            except PermanentIngestionError:
                continue
            for source_column, sector_code, sector_name in sectors:
                rows.append(
                    OpenCedaParameterRow(
                        release_year=release_year,
                        base_year=base_year,
                        parameter_type="sector_price_index",
                        name=f"Sector price index — {sector_name}",
                        value=self._decimal(
                            values[source_column - 1],
                            sheet.title,
                            source_row,
                            source_column,
                        ),
                        unit=f"index_{base_year}=100",
                        reference_year=reference_year,
                        sector_code=sector_code,
                        sector_name=sector_name,
                        source_sheet=sheet.title,
                        source_row=source_row,
                        source_column=source_column,
                    )
                )
        return tuple(rows)

    def _country_to_region(self, sheet: Worksheet) -> dict[str, tuple[str, ...]]:
        by_region: dict[str, list[str]] = {}
        for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
            country_code, region_name = row
            if country_code is None or region_name is None:
                continue
            by_region.setdefault(str(region_name).strip(), []).append(str(country_code).strip())
        return {key: tuple(value) for key, value in by_region.items()}

    def _sectors(
        self, sheet: Worksheet, name_row: int, code_row: int, first_column: int
    ) -> tuple[tuple[int, str, str], ...]:
        sectors: list[tuple[int, str, str]] = []
        maximum_column = first_column + self.expected_sector_count
        names = next(
            sheet.iter_rows(
                min_row=name_row,
                max_row=name_row,
                max_col=maximum_column,
                values_only=True,
            )
        )
        codes = next(
            sheet.iter_rows(
                min_row=code_row,
                max_row=code_row,
                max_col=maximum_column,
                values_only=True,
            )
        )
        for source_column in range(first_column, maximum_column):
            code = codes[source_column - 1]
            name = names[source_column - 1]
            if code is None or name is None:
                raise PermanentIngestionError(
                    f"Open CEDA sector schema changed at {sheet.title} column {source_column}"
                )
            sectors.append((source_column, str(code).strip(), str(name).strip()))
        next_code = codes[maximum_column - 1]
        if next_code is not None:
            raise PermanentIngestionError(
                f"Open CEDA sector count exceeds {self.expected_sector_count} in {sheet.title}"
            )
        return tuple(sectors)

    @staticmethod
    def _require_sheets(workbook: object, expected: set[str]) -> None:
        actual = set(workbook.sheetnames)  # type: ignore[attr-defined]
        missing = expected - actual
        if missing:
            raise PermanentIngestionError(f"Open CEDA workbook sheets changed: {sorted(missing)}")

    @staticmethod
    def _decimal(value: object, sheet: str, row: int, column: int) -> Decimal:
        if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
            raise PermanentIngestionError(
                f"Open CEDA numeric value is invalid at {sheet}!R{row}C{column}"
            )
        return Decimal(str(value))

    @staticmethod
    def _year(value: object, label: str) -> int:
        try:
            year = int(str(value))
        except ValueError as error:
            raise PermanentIngestionError(f"Open CEDA {label} is invalid: {value!r}") from error
        if not 1900 <= year <= 2200:
            raise PermanentIngestionError(f"Open CEDA {label} is invalid: {year}")
        return year

    @staticmethod
    def _published_at(value: object) -> datetime:
        if isinstance(value, datetime):
            return value.replace(tzinfo=value.tzinfo or UTC)
        normalized = re.sub(r"(\d+)(?:st|nd|rd|th)", r"\1", str(value)).strip()
        try:
            return datetime.strptime(normalized, "%B %d, %Y").replace(tzinfo=UTC)
        except ValueError as error:
            raise PermanentIngestionError(
                f"Open CEDA release date is invalid: {value!r}"
            ) from error

    @staticmethod
    def _slug(value: str) -> str:
        return re.sub(r"[^A-Z0-9]+", "_", value.upper()).strip("_")
