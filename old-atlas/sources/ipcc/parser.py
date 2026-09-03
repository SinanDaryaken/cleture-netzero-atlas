from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError

EXPECTED_HEADERS = (
    "EF ID",
    "IPCC 1996 Source/Sink Category",
    "IPCC 2006 Source/Sink Category",
    "Gas",
    "Fuel 1996",
    "Fuel 2006",
    "C pool",
    "Type of parameter",
    "Description",
    "Technologies / Practices",
    "Parameters / Conditions",
    "Region / Regional Conditions",
    "Abatement / Control Technologies",
    "Other properties",
    "Value",
    "Unit",
    "Equation",
    "IPCC Worksheet",
    "Technical Reference",
    "Source of data",
    "Data provider",
)


class IpccRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    ef_id: str
    category_1996: str | None = None
    category_2006: str | None = None
    gas: str | None = None
    fuel_1996: str | None = None
    fuel_2006: str | None = None
    carbon_pool: str | None = None
    parameter_type: str
    description: str | None = None
    technology: str | None = None
    conditions: str | None = None
    regional_conditions: str | None = None
    abatement: str | None = None
    other_properties: str | None = None
    raw_value: str | None = None
    value: Decimal | None = None
    unit: str | None = None
    equation: str | None = None
    worksheet: str | None = None
    technical_reference: str | None = None
    source_of_data: str | None = None
    data_provider: str | None = None
    source_row: int

    @property
    def is_default(self) -> bool:
        return "default" in self.parameter_type.lower()

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet="Sheet1",
            source_table="IPCC EFDB export",
        )


class IpccParser:
    """Parse the 21-column workbook exported by the official EFDB web application."""

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(self, path: Path) -> tuple[IpccRow, ...]:
        # EFDB labels the OOXML response as XLS. Passing a binary stream lets
        # openpyxl inspect the package signature instead of rejecting the suffix.
        with path.open("rb") as workbook_file:
            workbook = load_workbook(workbook_file, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if sheet is None:
                    raise PermanentIngestionError("IPCC EFDB export has no worksheet")
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    raise PermanentIngestionError("IPCC EFDB export is empty")
                actual_headers = tuple(self._text(value) or "" for value in header)
                if actual_headers != EXPECTED_HEADERS:
                    raise PermanentIngestionError(
                        "IPCC EFDB export schema changed; parser version must be updated"
                    )

                parsed: list[IpccRow] = []
                non_numeric = 0
                missing_value = 0
                for source_row, values in enumerate(rows, start=2):
                    row = self._row(values, source_row)
                    if row is None:
                        continue
                    parsed.append(row)
                    if row.raw_value is None:
                        missing_value += 1
                    elif row.value is None:
                        non_numeric += 1
            finally:
                workbook.close()

        if not parsed:
            raise PermanentIngestionError("IPCC EFDB export contains no records")
        self.metrics = {
            "exported_rows": len(parsed),
            "numeric_values": sum(row.value is not None for row in parsed),
            "non_numeric_values": non_numeric,
            "missing_values": missing_value,
            "default_rows": sum(row.is_default for row in parsed),
        }
        return tuple(parsed)

    def _row(self, values: tuple[Any, ...], source_row: int) -> IpccRow | None:
        padded = (*values, *(None for _ in range(max(0, len(EXPECTED_HEADERS) - len(values)))))
        ef_id = self._text(padded[0])
        parameter_type = self._text(padded[7])
        description = self._text(padded[8])
        if not ef_id and not description:
            return None
        if not ef_id or not parameter_type:
            raise PermanentIngestionError(
                f"IPCC EFDB mandatory field is missing at row {source_row}"
            )
        raw_value = self._text(padded[14])
        return IpccRow(
            ef_id=ef_id,
            category_1996=self._text(padded[1]),
            category_2006=self._text(padded[2]),
            gas=self._text(padded[3]),
            fuel_1996=self._text(padded[4]),
            fuel_2006=self._text(padded[5]),
            carbon_pool=self._text(padded[6]),
            parameter_type=parameter_type,
            description=description,
            technology=self._text(padded[9]),
            conditions=self._text(padded[10]),
            regional_conditions=self._text(padded[11]),
            abatement=self._text(padded[12]),
            other_properties=self._text(padded[13]),
            raw_value=raw_value,
            value=self._decimal(raw_value),
            unit=self._text(padded[15]),
            equation=self._text(padded[16]),
            worksheet=self._text(padded[17]),
            technical_reference=self._text(padded[18]),
            source_of_data=self._text(padded[19]),
            data_provider=self._text(padded[20]),
            source_row=source_row,
        )

    @staticmethod
    def _text(value: Any) -> str | None:
        if value is None:
            return None
        normalized = " ".join(str(value).split())
        return normalized or None

    @staticmethod
    def _decimal(value: str | None) -> Decimal | None:
        if value is None:
            return None
        try:
            return Decimal(value.replace(",", ""))
        except InvalidOperation:
            return None
