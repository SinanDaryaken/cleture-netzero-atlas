from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, ClassVar, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError


class CbamDefaultValueRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    row_kind: Literal["default_value"] = "default_value"
    geography_kind: Literal["country", "other", "annex_iv"]
    geography_name: str
    sector: str
    cn_code: str
    description: str
    direct_value: Decimal | None = None
    indirect_value: Decimal | None = None
    total_value: Decimal
    production_route: str | None = None
    original_file: str
    original_url: str
    member_sha256: str
    source_sheet: str
    source_row: int

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet=self.source_sheet,
            source_table="default_values",
        )


class CbamBenchmarkRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    row_kind: Literal["benchmark"] = "benchmark"
    sector: str
    cn_code: str
    description: str
    benchmark_column: Literal["A", "B"]
    benchmark_kind: Literal["process_related", "default"]
    value: Decimal
    production_route: str | None = None
    valid_from_year: int
    valid_to_year: int
    original_file: str
    original_url: str
    member_sha256: str
    source_sheet: str
    source_row: int
    source_column: int

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet=self.source_sheet,
            source_table="benchmarks",
        )


class CbamDocument(BaseModel):
    model_config = ConfigDict(frozen=True)

    default_values: tuple[CbamDefaultValueRow, ...]
    benchmarks: tuple[CbamBenchmarkRow, ...]

    @property
    def parsed_records(self) -> tuple[ParsedRecord, ...]:
        return tuple(
            [row.as_parsed_record() for row in self.default_values]
            + [row.as_parsed_record() for row in self.benchmarks]
        )


class CbamParser:
    expected_dataset_version = "definitive-2026-v2"
    _SECTORS: ClassVar[set[str]] = {
        "Cement",
        "Fertilisers",
        "Iron and steel",
        "Iron & Steel",
        "Aluminium",
        "Hydrogen",
    }

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(self, bundle_path: Path) -> CbamDocument:
        default_values: list[CbamDefaultValueRow] = []
        benchmarks: list[CbamBenchmarkRow] = []
        excluded_default_rows = 0
        with zipfile.ZipFile(bundle_path) as bundle:
            manifest = self._manifest(bundle)
            assets = {str(asset["role"]): asset for asset in manifest["assets"]}
            if set(assets) != {"default_values", "benchmarks"}:
                raise PermanentIngestionError("CBAM bundle member roles changed")
            for role, asset in assets.items():
                filename = str(asset["filename"])
                content = bundle.read(filename)
                checksum = hashlib.sha256(content).hexdigest()
                if checksum != asset["sha256"]:
                    raise PermanentIngestionError(
                        f"CBAM bundle member checksum mismatch: {filename}"
                    )
                if role == "default_values":
                    parsed, excluded_default_rows = self._parse_default_values(
                        content,
                        filename=filename,
                        source_url=str(asset["source_url"]),
                        member_sha256=checksum,
                    )
                    default_values.extend(parsed)
                else:
                    benchmarks.extend(
                        self._parse_benchmarks(
                            content,
                            filename=filename,
                            source_url=str(asset["source_url"]),
                            member_sha256=checksum,
                        )
                    )
        if not default_values or not benchmarks:
            raise PermanentIngestionError("CBAM workbooks produced incomplete source data")
        sectors = Counter(row.sector for row in default_values)
        geography_kinds = Counter(row.geography_kind for row in default_values)
        benchmark_columns = Counter(row.benchmark_column for row in benchmarks)
        self.metrics = {
            "parsed_rows": len(default_values) + len(benchmarks),
            "parsed_default_values": len(default_values),
            "parsed_benchmarks": len(benchmarks),
            "excluded_default_rows": excluded_default_rows,
            "country_default_values": geography_kinds["country"],
            "other_country_default_values": geography_kinds["other"],
            "annex_iv_default_values": geography_kinds["annex_iv"],
            "benchmark_column_a": benchmark_columns["A"],
            "benchmark_column_b": benchmark_columns["B"],
            **{
                f"table_{self._sector_key(sector)}_factors": count
                for sector, count in sectors.items()
            },
        }
        return CbamDocument(default_values=tuple(default_values), benchmarks=tuple(benchmarks))

    def _parse_default_values(
        self,
        content: bytes,
        *,
        filename: str,
        source_url: str,
        member_sha256: str,
    ) -> tuple[list[CbamDefaultValueRow], int]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise PermanentIngestionError("CBAM default-values workbook is unreadable") from error
        rows: list[CbamDefaultValueRow] = []
        excluded = 0
        try:
            required = {"Overview", "Version History", "Annex IV"}
            if not required.issubset(workbook.sheetnames):
                raise PermanentIngestionError("CBAM default-values workbook schema changed")
            overview = self._worksheet_text(workbook["Overview"])
            history = self._worksheet_text(workbook["Version History"])
            if (
                "2025/2621" not in overview
                or "2026/1740" not in history
                or "2026-08-06" not in history
            ):
                raise PermanentIngestionError("CBAM corrected default-values version changed")
            for sheet in workbook.worksheets:
                if sheet.title in {"Overview", "Version History"}:
                    continue
                geography_name = self._text(sheet.cell(1, 1).value)
                if geography_name is None:
                    raise PermanentIngestionError(f"CBAM geography name missing: {sheet.title}")
                if sheet.title == "Annex IV":
                    geography_kind: Literal["country", "other", "annex_iv"] = "annex_iv"
                elif geography_name == "Other Countries and Territories":
                    geography_kind = "other"
                else:
                    geography_kind = "country"
                sector: str | None = None
                for source_row, values in enumerate(
                    sheet.iter_rows(min_row=3, values_only=True), start=3
                ):
                    first = self._text(values[0] if values else None)
                    if first in self._SECTORS:
                        sector = first
                        continue
                    description = self._text(values[1] if len(values) > 1 else None)
                    cn_code = self._cn_code(first)
                    if sector is None or cn_code is None or description is None:
                        continue
                    if geography_kind == "annex_iv":
                        total = self._decimal(values[2] if len(values) > 2 else None)
                        direct = indirect = None
                        route = self._text(values[3] if len(values) > 3 else None)
                    else:
                        direct = self._decimal(values[2] if len(values) > 2 else None)
                        indirect = self._decimal(values[3] if len(values) > 3 else None)
                        total = self._decimal(values[4] if len(values) > 4 else None)
                        route = self._text(values[5] if len(values) > 5 else None)
                    if total is None:
                        excluded += 1
                        continue
                    rows.append(
                        CbamDefaultValueRow(
                            geography_kind=geography_kind,
                            geography_name=geography_name,
                            sector=sector,
                            cn_code=cn_code,
                            description=description,
                            direct_value=direct,
                            indirect_value=indirect,
                            total_value=total,
                            production_route=route,
                            original_file=filename,
                            original_url=source_url,
                            member_sha256=member_sha256,
                            source_sheet=sheet.title,
                            source_row=source_row,
                        )
                    )
        finally:
            workbook.close()
        return rows, excluded

    def _parse_benchmarks(
        self,
        content: bytes,
        *,
        filename: str,
        source_url: str,
        member_sha256: str,
    ) -> list[CbamBenchmarkRow]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise PermanentIngestionError("CBAM benchmark workbook is unreadable") from error
        rows: list[CbamBenchmarkRow] = []
        try:
            required = {"Overview", "Version history", "Benchmarks"}
            if not required.issubset(workbook.sheetnames):
                raise PermanentIngestionError("CBAM benchmark workbook schema changed")
            if "2025/2620" not in self._worksheet_text(workbook["Overview"]):
                raise PermanentIngestionError("CBAM benchmark legal basis changed")
            sheet = workbook["Benchmarks"]
            sector: str | None = None
            cn_code: str | None = None
            description: str | None = None
            for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                first = self._text(values[0] if values else None)
                if first in self._SECTORS:
                    sector = first
                    cn_code = description = None
                    continue
                next_cn = self._cn_code(first)
                next_description = self._text(values[1] if len(values) > 1 else None)
                if next_cn is not None:
                    cn_code = next_cn
                    description = next_description
                if sector is None or cn_code is None or description is None:
                    continue
                for column, value_index, route_index, kind in (
                    ("A", 2, 3, "process_related"),
                    ("B", 4, 5, "default"),
                ):
                    value = self._decimal(
                        values[value_index] if len(values) > value_index else None
                    )
                    if value is None:
                        continue
                    route = self._text(values[route_index] if len(values) > route_index else None)
                    valid_from, valid_to = self._validity(route)
                    rows.append(
                        CbamBenchmarkRow(
                            sector=sector,
                            cn_code=cn_code,
                            description=description,
                            benchmark_column=column,
                            benchmark_kind=kind,
                            value=value,
                            production_route=route,
                            valid_from_year=valid_from,
                            valid_to_year=valid_to,
                            original_file=filename,
                            original_url=source_url,
                            member_sha256=member_sha256,
                            source_sheet=sheet.title,
                            source_row=source_row,
                            source_column=value_index + 1,
                        )
                    )
        finally:
            workbook.close()
        return rows

    def _manifest(self, bundle: zipfile.ZipFile) -> dict[str, Any]:
        try:
            payload = json.loads(bundle.read("bundle-manifest.json"))
        except (KeyError, json.JSONDecodeError, UnicodeDecodeError) as error:
            raise PermanentIngestionError("CBAM bundle manifest is invalid") from error
        if (
            not isinstance(payload, dict)
            or payload.get("dataset_version") != self.expected_dataset_version
        ):
            raise PermanentIngestionError("CBAM bundle version changed")
        assets = payload.get("assets")
        if not isinstance(assets, list) or not all(isinstance(asset, dict) for asset in assets):
            raise PermanentIngestionError("CBAM bundle assets are missing")
        return payload

    @staticmethod
    def _worksheet_text(sheet: Any) -> str:
        return " ".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )

    @classmethod
    def _cn_code(cls, value: object) -> str | None:
        text = cls._text(value)
        if text is None:
            return None
        digits = re.sub(r"\D", "", text)
        return digits if 4 <= len(digits) <= 10 else None

    @staticmethod
    def _validity(route: str | None) -> tuple[int, int]:
        if route and "(1)" in route:
            return 2026, 2027
        if route and "(2)" in route:
            return 2028, 2030
        return 2026, 2030

    @staticmethod
    def _sector_key(sector: str) -> str:
        return {
            "Cement": "cement",
            "Fertilisers": "fertilisers",
            "Iron and steel": "iron_steel",
            "Iron & Steel": "iron_steel",
            "Aluminium": "aluminium",
            "Hydrogen": "hydrogen",
        }[sector]

    @staticmethod
    def _text(value: object) -> str | None:
        if value is None:
            return None
        text = " ".join(str(value).split())
        return text or None

    @staticmethod
    def _decimal(value: object) -> Decimal | None:
        if value is None or isinstance(value, bool):
            return None
        text = str(value).strip().replace("\u00a0", "").replace(",", ".")
        try:
            return Decimal(text)
        except (InvalidOperation, ValueError):
            return None
