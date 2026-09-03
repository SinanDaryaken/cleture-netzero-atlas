from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, ClassVar

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from atlas.domain.enums import EnvironmentalEntityType
from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError


class UnfcccTuikRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    submission_year: int
    submission_status: str
    reference_year: int
    schema_family: str
    entity_type: EnvironmentalEntityType
    category: str
    category_path: str
    column_label: str
    gas: str | None = None
    value: Decimal
    unit: str
    activity_unit: str | None = None
    member_filename: str
    member_sha256: str
    source_sheet: str
    source_row: int
    source_column: int

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet=self.source_sheet,
            source_table=self.schema_family,
        )


class UnfcccTuikParser:
    """Parse numeric CRF/CRT background-table observations without flattening semantics."""

    max_template_columns = 64
    header_scan_rows = 15
    header_detail_rows = 3
    notation_keys: ClassVar[set[str]] = {"C", "CR", "IE", "NA", "NE", "NO"}

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(
        self,
        archive_path: Path,
        *,
        submission_year: int,
        submission_status: str,
        schema_family: str,
    ) -> tuple[UnfcccTuikRow, ...]:
        rows: list[UnfcccTuikRow] = []
        members = 0
        with zipfile.ZipFile(archive_path) as archive:
            workbook_members = sorted(
                name for name in archive.namelist() if name.lower().endswith(".xlsx")
            )
            if not workbook_members:
                raise PermanentIngestionError("UNFCCC submission contains no XLSX workbooks")
            for member in workbook_members:
                reference_year = self._reference_year(member, submission_year)
                content = archive.read(member)
                member_sha256 = hashlib.sha256(content).hexdigest()
                rows.extend(
                    self._parse_workbook(
                        content,
                        submission_year=submission_year,
                        submission_status=submission_status,
                        reference_year=reference_year,
                        schema_family=schema_family,
                        member_filename=member,
                        member_sha256=member_sha256,
                    )
                )
                members += 1
        if not rows:
            raise PermanentIngestionError(
                "UNFCCC submission produced no structured CRF/CRT observations"
            )
        entity_counts = Counter(row.entity_type.value for row in rows)
        reference_years = {row.reference_year for row in rows}
        self.metrics = {
            "parsed_rows": len(rows),
            "workbook_members": members,
            "reference_year_min": min(reference_years),
            "reference_year_max": max(reference_years),
            **{f"entity_{key}": value for key, value in entity_counts.items()},
        }
        return tuple(rows)

    def _parse_workbook(
        self,
        content: bytes,
        *,
        submission_year: int,
        submission_status: str,
        reference_year: int,
        schema_family: str,
        member_filename: str,
        member_sha256: str,
    ) -> list[UnfcccTuikRow]:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[UnfcccTuikRow] = []
        try:
            for sheet in workbook.worksheets:
                matrix = list(
                    sheet.iter_rows(
                        max_col=min(sheet.max_column, self.max_template_columns),
                        values_only=True,
                    )
                )
                header = self._header(matrix)
                if header is None:
                    continue
                header_row, category_column, groups = header
                data_start_row = self._data_start(matrix, header_row, category_column, groups)
                activity_unit_column = self._activity_unit_column(
                    matrix,
                    header_row,
                    data_start_row,
                    groups,
                )
                last_coded_category: str | None = None
                for source_row, values in enumerate(
                    matrix[data_start_row - 1 :],
                    start=data_start_row,
                ):
                    category = self._text(values[category_column - 1])
                    if not self._is_data_category(category):
                        continue
                    assert category is not None
                    if self._is_coded_category(category):
                        last_coded_category = category
                    category_path = (
                        category
                        if last_coded_category is None or last_coded_category == category
                        else f"{last_coded_category} / {category}"
                    )
                    row_activity_unit = (
                        self._text(values[activity_unit_column - 1])
                        if activity_unit_column is not None
                        else None
                    )
                    for entity_type, start_column, end_column in groups:
                        group_unit = self._group_unit(
                            matrix, header_row, data_start_row, start_column, end_column
                        )
                        for column in range(start_column, end_column + 1):
                            value = self._decimal(values[column - 1])
                            if value is None:
                                continue
                            label, explicit_unit = self._column_descriptor(
                                matrix, header_row, data_start_row, column
                            )
                            unit = explicit_unit or group_unit or "source_unit_unspecified"
                            gas = self._gas(label)
                            unit = self._repair_inherited_ief_unit(
                                entity_type=entity_type,
                                gas=gas,
                                unit=unit,
                            )
                            rows.append(
                                UnfcccTuikRow(
                                    submission_year=submission_year,
                                    submission_status=submission_status,
                                    reference_year=reference_year,
                                    schema_family=schema_family,
                                    entity_type=entity_type,
                                    category=category,
                                    category_path=category_path,
                                    column_label=label,
                                    gas=gas,
                                    value=value,
                                    unit=unit,
                                    activity_unit=row_activity_unit,
                                    member_filename=member_filename,
                                    member_sha256=member_sha256,
                                    source_sheet=sheet.title,
                                    source_row=source_row,
                                    source_column=column,
                                )
                            )
        finally:
            workbook.close()
        return rows

    def _activity_unit_column(
        self,
        matrix: list[tuple[Any, ...]],
        header_row: int,
        data_start_row: int,
        groups: list[tuple[EnvironmentalEntityType, int, int]],
    ) -> int | None:
        for entity_type, start_column, end_column in groups:
            if entity_type != EnvironmentalEntityType.ACTIVITY_DATA:
                continue
            for row in range(header_row + 1, data_start_row):
                for column in range(start_column, end_column + 1):
                    text = self._text(matrix[row - 1][column - 1])
                    if text and text.casefold().strip().startswith("unit"):
                        return column
        return None

    @staticmethod
    def _repair_inherited_ief_unit(
        *,
        entity_type: EnvironmentalEntityType,
        gas: str | None,
        unit: str,
    ) -> str:
        # CRT energy sheets leave the N2O unit cell blank because CH4 and N2O
        # share kg/TJ. The generic group fallback otherwise inherits CO2 t/TJ.
        if (
            entity_type == EnvironmentalEntityType.IMPLIED_EMISSION_FACTOR
            and gas in {"CH4", "N2O"}
            and unit == "t/TJ"
        ):
            return "kg/TJ"
        return unit

    def _header(
        self, matrix: list[tuple[Any, ...]]
    ) -> tuple[int, int, list[tuple[EnvironmentalEntityType, int, int]]] | None:
        max_column = len(matrix[0]) if matrix else 0
        for row_number, values in enumerate(matrix[: self.header_scan_rows], start=1):
            texts = [self._text(value) for value in values]
            implied_columns = [
                index
                for index, text in enumerate(texts, start=1)
                if text and "IMPLIED EMISSION FACTOR" in text.upper()
            ]
            if not implied_columns:
                continue
            starts = [(index, text) for index, text in enumerate(texts, start=1) if text]
            first_measure = min(
                index for index, text in starts if self._entity_type(text) is not None
            )
            category_column = max(
                (index for index, _ in starts if index < first_measure),
                default=max(first_measure - 1, 1),
            )
            groups: list[tuple[EnvironmentalEntityType, int, int]] = []
            for position, (start_column, text) in enumerate(starts):
                entity_type = self._entity_type(text)
                if entity_type is None:
                    continue
                next_start = (
                    starts[position + 1][0] if position + 1 < len(starts) else max_column + 1
                )
                groups.append((entity_type, start_column, next_start - 1))
            return row_number, category_column, groups
        return None

    @staticmethod
    def _entity_type(text: str) -> EnvironmentalEntityType | None:
        normalized = " ".join(text.upper().split())
        if "IMPLIED EMISSION FACTOR" in normalized:
            return EnvironmentalEntityType.IMPLIED_EMISSION_FACTOR
        if "ACTIVITY DATA" in normalized:
            return EnvironmentalEntityType.ACTIVITY_DATA
        if normalized.startswith("EMISSIONS") or normalized == "EMISSION":
            return EnvironmentalEntityType.EMISSION_RESULT
        return None

    def _data_start(
        self,
        matrix: list[tuple[Any, ...]],
        header_row: int,
        category_column: int,
        groups: list[tuple[EnvironmentalEntityType, int, int]],
    ) -> int:
        measure_columns = [column for _, start, end in groups for column in range(start, end + 1)]
        for row_number in range(header_row + 1, min(len(matrix), header_row + 8) + 1):
            values = matrix[row_number - 1]
            category = self._text(values[category_column - 1])
            if not self._is_data_category(category):
                continue
            if any(self._decimal(values[column - 1]) is not None for column in measure_columns):
                return row_number
        return header_row + self.header_detail_rows + 1

    def _column_descriptor(
        self,
        matrix: list[tuple[Any, ...]],
        header_row: int,
        data_start_row: int,
        column: int,
    ) -> tuple[str, str | None]:
        labels: list[str] = []
        unit: str | None = None
        for row in range(header_row + 1, data_start_row):
            text = self._text(matrix[row - 1][column - 1])
            if text is None:
                continue
            if self._looks_like_unit(text):
                unit = self._normalize_unit(text)
            else:
                labels.append(" ".join(text.split()))
        return " / ".join(labels) or f"column_{column}", unit

    def _group_unit(
        self,
        matrix: list[tuple[Any, ...]],
        header_row: int,
        data_start_row: int,
        start_column: int,
        end_column: int,
    ) -> str | None:
        for row in range(header_row + 1, data_start_row):
            for column in range(start_column, end_column + 1):
                text = self._text(matrix[row - 1][column - 1])
                if text and self._looks_like_unit(text):
                    return self._normalize_unit(text)
        return None

    @staticmethod
    def _looks_like_unit(text: str) -> bool:
        normalized = " ".join(text.split())
        return (normalized.startswith("(") and normalized.endswith(")")) or normalized.lower() in {
            "kg",
            "kt",
            "t",
            "tj",
            "%",
            "ha",
        }

    @staticmethod
    def _normalize_unit(text: str) -> str:
        return " ".join(text.strip().strip("()").split())

    @classmethod
    def _decimal(cls, value: Any) -> Decimal | None:
        if value is None or isinstance(value, bool):
            return None
        if isinstance(value, Decimal):
            return value
        if isinstance(value, (int, float)):
            return Decimal(str(value))
        text = str(value).strip()
        if not text or set(part.strip() for part in text.split(",")) <= cls.notation_keys:
            return None
        try:
            return Decimal(text.replace(" ", "").replace(",", "."))
        except InvalidOperation:
            return None

    @staticmethod
    def _text(value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _is_data_category(category: str | None) -> bool:
        if category is None:
            return False
        normalized = category.strip().lower()
        return not (
            normalized.startswith(("(", "note", "documentation", "source", "back to"))
            or normalized.startswith("•")
            or "parties should" in normalized
        )

    @staticmethod
    def _is_coded_category(category: str) -> bool:
        return bool(re.match(r"^(?:[1-5](?:\.|\s)|[A-E]\.)", category.strip()))

    @staticmethod
    def _gas(label: str) -> str | None:
        normalized = label.upper().replace("_X000D_", " ")
        for gas in ("CO2", "CH4", "N2O", "SF6", "NF3", "NOX", "NMVOC", "SOX", "CO"):
            if re.search(rf"\b{gas}\b", normalized):
                return gas
        return None

    @staticmethod
    def _reference_year(member_filename: str, submission_year: int) -> int:
        patterns = (
            rf"_{submission_year}_((?:19|20)\d{{2}})_",
            rf"-CRT-{submission_year}-[^-]+-((?:19|20)\d{{2}})-",
        )
        for pattern in patterns:
            match = re.search(pattern, member_filename)
            if match:
                return int(match.group(1))
        years = [int(value) for value in re.findall(r"(?:19|20)\d{2}", member_filename)]
        candidates = [year for year in years if year <= submission_year]
        if not candidates:
            raise PermanentIngestionError(
                f"cannot derive reference year from UNFCCC member: {member_filename}"
            )
        return candidates[-1]
