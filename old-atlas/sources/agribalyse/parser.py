from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError


class AgribalyseRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    dataset_part: Literal["conventional", "organic", "food"]
    source_code: str
    name_fr: str
    lci_name: str
    group: str | None = None
    subgroup: str | None = None
    production_type: str | None = None
    source_category: str | None = None
    season_code: str | None = None
    air_freight_code: str | None = None
    delivery: str | None = None
    packaging: str | None = None
    preparation: str | None = None
    dqr: Decimal | None = None
    climate_change_kgco2e_per_kg: Decimal
    geography_code: str
    duplicate_index: int
    original_file: str
    original_url: str
    member_sha256: str
    source_sheet: str
    source_row: int

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet=self.source_sheet,
            source_table=self.dataset_part,
        )


class AgribalyseParser:
    expected_version = "3.2"

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(self, bundle_path: Path) -> tuple[AgribalyseRow, ...]:
        rows: list[AgribalyseRow] = []
        with zipfile.ZipFile(bundle_path) as bundle:
            manifest = self._manifest(bundle)
            for asset in manifest["assets"]:
                if not isinstance(asset, dict):
                    raise PermanentIngestionError("AGRIBALYSE bundle asset metadata changed")
                role = str(asset["role"])
                filename = str(asset["filename"])
                content = bundle.read(filename)
                checksum = hashlib.sha256(content).hexdigest()
                if checksum != asset["sha256"]:
                    raise PermanentIngestionError(
                        f"AGRIBALYSE bundle member checksum mismatch: {filename}"
                    )
                rows.extend(
                    self._parse_workbook(
                        content,
                        role=role,
                        filename=filename,
                        source_url=str(asset["source_url"]),
                        member_sha256=checksum,
                    )
                )
        counts = Counter(row.dataset_part for row in rows)
        self.metrics = {
            "parsed_rows": len(rows),
            "conventional_rows": counts["conventional"],
            "organic_rows": counts["organic"],
            "food_rows": counts["food"],
            "negative_lca_results": sum(row.climate_change_kgco2e_per_kg < 0 for row in rows),
        }
        if not rows:
            raise PermanentIngestionError("AGRIBALYSE workbooks produced no LCA result rows")
        return tuple(rows)

    def _parse_workbook(
        self,
        content: bytes,
        *,
        role: str,
        filename: str,
        source_url: str,
        member_sha256: str,
    ) -> list[AgribalyseRow]:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            notice = " ".join(
                str(cell.value)
                for row in workbook["Notice"].iter_rows()
                for cell in row
                if cell.value
            )
            if "3.2" not in notice:
                raise PermanentIngestionError(f"AGRIBALYSE dataset version changed in {filename}")
            if role == "food":
                return self._food_rows(workbook["Synthese"], filename, source_url, member_sha256)
            if role == "conventional":
                agriculture_role: Literal["conventional", "organic"] = "conventional"
            elif role == "organic":
                agriculture_role = "organic"
            else:
                raise PermanentIngestionError(f"unknown AGRIBALYSE bundle role: {role}")
            sheet = next(
                sheet for sheet in workbook.worksheets if sheet.title.startswith("AGB 3.2")
            )
            return self._agriculture_rows(
                sheet, agriculture_role, filename, source_url, member_sha256
            )
        finally:
            workbook.close()

    def _agriculture_rows(
        self,
        sheet: Any,
        role: Literal["conventional", "organic"],
        filename: str,
        source_url: str,
        member_sha256: str,
    ) -> list[AgribalyseRow]:
        rows: list[AgribalyseRow] = []
        identities: Counter[tuple[str, ...]] = Counter()
        for source_row, values in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
            name = self._text(values[0])
            lci_name = self._text(values[1])
            climate = self._decimal(values[5])
            if name is None or lci_name is None or climate is None:
                continue
            identity = (role, name, lci_name)
            identities[identity] += 1
            rows.append(
                AgribalyseRow(
                    dataset_part=role,
                    source_code=lci_name,
                    name_fr=name,
                    lci_name=lci_name,
                    production_type=self._text(values[2]),
                    source_category=self._text(values[3]),
                    climate_change_kgco2e_per_kg=climate,
                    geography_code=self._location(lci_name),
                    duplicate_index=identities[identity],
                    original_file=filename,
                    original_url=source_url,
                    member_sha256=member_sha256,
                    source_sheet=sheet.title,
                    source_row=source_row,
                )
            )
        return rows

    def _food_rows(
        self,
        sheet: Any,
        filename: str,
        source_url: str,
        member_sha256: str,
    ) -> list[AgribalyseRow]:
        rows: list[AgribalyseRow] = []
        identities: Counter[tuple[str, ...]] = Counter()
        for source_row, values in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
            code = self._text(values[0])
            name = self._text(values[4])
            lci_name = self._text(values[5])
            climate = self._decimal(values[13])
            if code is None or name is None or lci_name is None or climate is None:
                continue
            scenario = tuple(self._text(values[index]) or "" for index in range(6, 11))
            identity = ("food", code, name, lci_name, *scenario)
            identities[identity] += 1
            rows.append(
                AgribalyseRow(
                    dataset_part="food",
                    source_code=code,
                    name_fr=name,
                    lci_name=lci_name,
                    group=self._text(values[2]),
                    subgroup=self._text(values[3]),
                    season_code=self._text(values[6]),
                    air_freight_code=self._text(values[7]),
                    delivery=self._text(values[8]),
                    packaging=self._text(values[9]),
                    preparation=self._text(values[10]),
                    dqr=self._decimal(values[11]),
                    climate_change_kgco2e_per_kg=climate,
                    geography_code="FR",
                    duplicate_index=identities[identity],
                    original_file=filename,
                    original_url=source_url,
                    member_sha256=member_sha256,
                    source_sheet=sheet.title,
                    source_row=source_row,
                )
            )
        return rows

    def _manifest(self, bundle: zipfile.ZipFile) -> dict[str, Any]:
        try:
            payload = json.loads(bundle.read("bundle-manifest.json"))
        except (KeyError, json.JSONDecodeError, UnicodeDecodeError) as error:
            raise PermanentIngestionError("AGRIBALYSE bundle manifest is invalid") from error
        if not isinstance(payload, dict) or payload.get("dataset_version") != self.expected_version:
            raise PermanentIngestionError("AGRIBALYSE bundle version changed")
        if not isinstance(payload.get("assets"), list):
            raise PermanentIngestionError("AGRIBALYSE bundle assets are missing")
        return payload

    @staticmethod
    def _location(lci_name: str) -> str:
        match = re.search(r"\{([^{}]+)\}", lci_name)
        return match.group(1) if match else "FR"

    @staticmethod
    def _text(value: object) -> str | None:
        if value is None:
            return None
        text = " ".join(str(value).split())
        return text or None

    @staticmethod
    def _decimal(value: object) -> Decimal | None:
        if value is None or isinstance(value, bool):
            return None
        try:
            return Decimal(str(value))
        except (ValueError, TypeError):
            return None
