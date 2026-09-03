from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import ClassVar

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError


class WrapRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    dataset_part: str
    product: str
    source_database: str
    source_database_id: str | None = None
    gpc_classification: str | None = None
    source_year: str | None = None
    origin_region: str
    applicable_region: str
    production_system: str | None = None
    intermediate_product: str | None = None
    emission_source: str | None = None
    lsr_category: str | None = None
    flag_category: str | None = None
    functional_unit: str
    lifecycle_stage: str
    factor_kg_co2e: Decimal
    data_quality_score: str | None = None
    standard_deviation: Decimal | None = None
    source_link: str | None = None
    source_row: int
    source_sheet: str

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet=self.source_sheet,
            source_table=self.dataset_part,
        )


class WrapParser:
    hestia_sheet = "emissions_database_hestia"
    refined_sheet = "emissions_database_refined"
    hestia_headers: ClassVar[set[str]] = {
        "source db",
        "product",
        "year",
        "country",
        "production technology",
        "functional unit",
        "production stage",
        "impact",
        "model",
        "unit",
        "value",
        "standard deviation",
        "aggregated quality",
        "link",
    }
    refined_headers: ClassVar[set[str]] = {
        "gpc_classification",
        "source_db_name",
        "source_db_id",
        "lifecycle_stage",
        "origin_region",
        "applicable_region",
        "production_system",
        "factor_kg_co2e",
        "func_unit",
        "factor_type",
        "data_quality_score",
    }

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(self, path: Path) -> tuple[WrapRow, ...]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            missing_sheets = {self.hestia_sheet, self.refined_sheet} - set(
                workbook.sheetnames
            )
            if missing_sheets:
                raise PermanentIngestionError(
                    "WRAP workbook schema changed; missing sheets: "
                    + ", ".join(sorted(missing_sheets))
                )
            rows = [*self._parse_hestia(workbook[self.hestia_sheet])]
            rows.extend(self._parse_refined(workbook[self.refined_sheet]))
        finally:
            workbook.close()
        if not rows:
            raise PermanentIngestionError("WRAP workbook produced no climate LCA results")
        self.metrics = {
            "input_rows": len(rows),
            "parsed_rows": len(rows),
            "hestia_gwp100_rows": sum(row.dataset_part == "hestia_gwp100" for row in rows),
            "refined_rows": sum(row.dataset_part == "refined" for row in rows),
            "negative_lca_results": sum(row.factor_kg_co2e < 0 for row in rows),
            "unique_products": len({row.product for row in rows}),
        }
        return tuple(rows)

    def _parse_hestia(self, sheet: object) -> list[WrapRow]:
        values = sheet.iter_rows(values_only=True)  # type: ignore[attr-defined]
        headers = tuple(self._text(value) or "" for value in next(values))
        self._require_headers(headers, self.hestia_headers, self.hestia_sheet)
        rows: list[WrapRow] = []
        for source_row, cells in enumerate(values, start=2):
            record = dict(zip(headers, cells, strict=False))
            if self._text(record.get("impact")) != "GWP100":
                continue
            if self._text(record.get("model")) != "IPCC2021":
                raise PermanentIngestionError(
                    f"WRAP HESTIA row {source_row} has an unexpected GWP model"
                )
            if self._text(record.get("unit")) != "kg CO2e":
                raise PermanentIngestionError(
                    f"WRAP HESTIA row {source_row} has an unexpected GWP unit"
                )
            country = self._required(record.get("country"), "country", source_row)
            rows.append(
                WrapRow(
                    dataset_part="hestia_gwp100",
                    product=self._required(record.get("product"), "product", source_row),
                    source_database=self._required(
                        record.get("source db"), "source db", source_row
                    ),
                    source_year=self._text(record.get("year")),
                    origin_region=country,
                    applicable_region=country,
                    production_system=self._text(record.get("production technology")),
                    intermediate_product=self._text(record.get("intermediate product")),
                    emission_source=self._text(record.get("emission source")),
                    lsr_category=self._text(record.get("LSR category")),
                    flag_category=self._text(record.get("FLAG category")),
                    functional_unit=self._required(
                        record.get("functional unit"), "functional unit", source_row
                    ),
                    lifecycle_stage=self._required(
                        record.get("production stage"), "production stage", source_row
                    ),
                    factor_kg_co2e=self._decimal(
                        record.get("value"), "value", source_row
                    ),
                    data_quality_score=self._text(record.get("aggregated quality")),
                    standard_deviation=self._optional_decimal(
                        record.get("standard deviation"), "standard deviation", source_row
                    ),
                    source_link=self._text(record.get("link")),
                    source_row=source_row,
                    source_sheet=self.hestia_sheet,
                )
            )
        return rows

    def _parse_refined(self, sheet: object) -> list[WrapRow]:
        values = sheet.iter_rows(values_only=True)  # type: ignore[attr-defined]
        headers = tuple(self._text(value) or "" for value in next(values))
        self._require_headers(headers, self.refined_headers, self.refined_sheet)
        rows: list[WrapRow] = []
        for source_row, cells in enumerate(values, start=2):
            record = dict(zip(headers, cells, strict=False))
            if not any(value is not None for value in cells):
                continue
            factor_type = self._required(record.get("factor_type"), "factor_type", source_row)
            if factor_type != "aggregate":
                raise PermanentIngestionError(
                    f"WRAP refined row {source_row} has an unexpected factor type"
                )
            rows.append(
                WrapRow(
                    dataset_part="refined",
                    product=self._required(
                        record.get("source_db_name"), "source_db_name", source_row
                    ),
                    source_database=self._required(
                        record.get("source_db_id"), "source_db_id", source_row
                    ),
                    source_database_id=self._required(
                        record.get("source_db_id"), "source_db_id", source_row
                    ),
                    gpc_classification=self._text(record.get("gpc_classification")),
                    origin_region=self._required(
                        record.get("origin_region"), "origin_region", source_row
                    ),
                    applicable_region=self._required(
                        record.get("applicable_region"), "applicable_region", source_row
                    ),
                    production_system=self._text(record.get("production_system")),
                    functional_unit=self._required(
                        record.get("func_unit"), "func_unit", source_row
                    ),
                    lifecycle_stage=self._required(
                        record.get("lifecycle_stage"), "lifecycle_stage", source_row
                    ),
                    factor_kg_co2e=self._decimal(
                        record.get("factor_kg_co2e"), "factor_kg_co2e", source_row
                    ),
                    data_quality_score=self._text(record.get("data_quality_score")),
                    source_row=source_row,
                    source_sheet=self.refined_sheet,
                )
            )
        return rows

    @staticmethod
    def _require_headers(headers: tuple[str, ...], expected: set[str], sheet: str) -> None:
        missing = expected - set(headers)
        if missing:
            raise PermanentIngestionError(
                f"WRAP workbook schema changed in {sheet}; missing headers: "
                + ", ".join(sorted(missing))
            )

    @staticmethod
    def _text(value: object) -> str | None:
        if value is None:
            return None
        text = " ".join(str(value).split())
        return text or None

    @classmethod
    def _required(cls, value: object, field: str, row: int) -> str:
        text = cls._text(value)
        if text is None:
            raise PermanentIngestionError(f"WRAP row {row} has no {field}")
        return text

    @classmethod
    def _decimal(cls, value: object, field: str, row: int) -> Decimal:
        result = cls._optional_decimal(value, field, row)
        if result is None:
            raise PermanentIngestionError(f"WRAP row {row} has no {field}")
        return result

    @staticmethod
    def _optional_decimal(value: object, field: str, row: int) -> Decimal | None:
        if value is None or not str(value).strip():
            return None
        try:
            result = Decimal(str(value))
        except InvalidOperation as error:
            raise PermanentIngestionError(
                f"WRAP row {row} has invalid {field}: {value}"
            ) from error
        if not result.is_finite():
            raise PermanentIngestionError(f"WRAP row {row} has non-finite {field}")
        return result
