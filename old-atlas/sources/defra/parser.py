from __future__ import annotations

import hashlib
from collections.abc import Iterable
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError

MODERN_REQUIRED_HEADERS = (
    "ID",
    "Scope",
    "Level 1",
    "Level 2",
    "Level 3",
    "Level 4",
    "Column Text",
    "UOM",
    "GHG/Unit",
    "GHG Conversion Factor",
)
LEGACY_REQUIRED_HEADERS = (
    "Scope",
    "Level 1",
    "Level 2",
    "Level 3",
    "Level 4",
    "Column Text",
    "UOM",
)


class DefraRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    source_id: str
    scope: str
    level_1: str
    level_2: str | None
    level_3: str | None
    level_4: str | None
    column_text: str | None
    uom: str
    ghg_unit: str
    value: Decimal | None
    row_number: int
    value_qualifier: str | None = None
    source_lookup: str | None = None

    @property
    def base_source_id(self) -> str:
        try:
            base, suffix = self.source_id.rsplit("_", 1)
        except ValueError as error:
            raise PermanentIngestionError(f"invalid DEFRA ID: {self.source_id}") from error
        if suffix not in {"1", "2", "3", "4", "5", "6"}:
            raise PermanentIngestionError(f"unknown DEFRA ID suffix: {self.source_id}")
        return base

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.row_number,
            source_sheet="Factors by Category",
        )

    @property
    def logical_key(self) -> str:
        identity = "\x1f".join(
            value or ""
            for value in (
                self.scope,
                self.level_1,
                self.level_2,
                self.level_3,
                self.level_4,
                self.column_text,
                self.uom,
            )
        )
        return hashlib.sha256(identity.casefold().encode()).hexdigest()[:24]


class DefraParser:
    def parse(self, path: Path) -> tuple[DefraRow, ...]:
        if path.suffix.lower() == ".xls":
            frame = pl.read_excel(
                path,
                sheet_name="Factors by Category",
                has_header=False,
                engine="calamine",
            )
            values = [tuple(row.values()) for row in frame.iter_rows(named=True)]
            return self._parse_rows(values)

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if "Factors by Category" not in workbook.sheetnames:
                raise PermanentIngestionError("DEFRA sheet 'Factors by Category' is missing")
            values = list(workbook["Factors by Category"].iter_rows(values_only=True))
            return self._parse_rows(values)
        finally:
            workbook.close()

    def _parse_rows(self, source_rows: list[tuple[Any, ...]]) -> tuple[DefraRow, ...]:
        header_row, headers, has_source_id = self._find_headers(source_rows[:20])
        rows: list[DefraRow] = []
        for index, values in enumerate(source_rows[header_row:], start=header_row + 1):
            payload = dict(zip(headers, values, strict=False))
            scope = self._text(payload.get("Scope"))
            level_1 = self._text(payload.get("Level 1"))
            uom = self._normalize_uom(self._text(payload.get("UOM")))
            raw_ghg_unit = self._text(payload.get("GHG/Unit") or payload.get("GHG"))
            if scope is None and level_1 is None and raw_ghg_unit is None:
                continue
            if scope is None or level_1 is None or uom is None or raw_ghg_unit is None:
                continue
            ghg_unit = self._normalize_ghg_unit(raw_ghg_unit)
            source_lookup = self._text(payload.get("Lookup"))
            source_id = self._text(payload.get("ID")) if has_source_id else None
            if source_id is None:
                source_id = self._legacy_source_id(
                    scope=scope,
                    level_1=level_1,
                    level_2=self._text(payload.get("Level 2")),
                    level_3=self._text(payload.get("Level 3")),
                    level_4=self._text(payload.get("Level 4")),
                    column_text=self._text(payload.get("Column Text")),
                    uom=uom,
                    source_lookup=source_lookup,
                    ghg_unit=ghg_unit,
                )
            raw_value = payload.get("GHG Conversion Factor")
            rows.append(
                DefraRow(
                    source_id=source_id,
                    scope=scope,
                    level_1=level_1,
                    level_2=self._text(payload.get("Level 2")),
                    level_3=self._text(payload.get("Level 3")),
                    level_4=self._text(payload.get("Level 4")),
                    column_text=self._text(payload.get("Column Text")),
                    uom=uom,
                    ghg_unit=ghg_unit,
                    value=self._decimal(raw_value, index),
                    value_qualifier=self._value_qualifier(raw_value),
                    row_number=index,
                    source_lookup=source_lookup,
                )
            )
        if not rows:
            raise PermanentIngestionError("DEFRA workbook contains no factor rows")
        if len({row.source_id for row in rows}) != len(rows):
            raise PermanentIngestionError("DEFRA workbook contains duplicate source IDs")
        return tuple(rows)

    @staticmethod
    def _find_headers(rows: Iterable[tuple[Any, ...]]) -> tuple[int, tuple[str, ...], bool]:
        for index, row in enumerate(rows, start=1):
            values = tuple(str(value).strip() if value is not None else "" for value in row)
            has_conversion_factor = any(
                value.startswith("GHG Conversion Factor ") for value in values
            )
            if not has_conversion_factor:
                continue
            has_source_id = set(MODERN_REQUIRED_HEADERS[:-1]).issubset(values)
            has_legacy_shape = set(LEGACY_REQUIRED_HEADERS).issubset(values) and (
                "GHG" in values or "GHG/Unit" in values
            )
            if has_source_id or has_legacy_shape:
                normalized = tuple(
                    "GHG Conversion Factor" if value.startswith("GHG Conversion Factor ") else value
                    for value in values
                )
                return index, normalized, has_source_id
        raise PermanentIngestionError("DEFRA header schema was not found")

    @classmethod
    def _legacy_source_id(
        cls,
        *,
        scope: str,
        level_1: str,
        level_2: str | None,
        level_3: str | None,
        level_4: str | None,
        column_text: str | None,
        uom: str,
        source_lookup: str | None,
        ghg_unit: str,
    ) -> str:
        identity = "\x1f".join(
            value or ""
            for value in (
                scope,
                level_1,
                level_2,
                level_3,
                level_4,
                column_text,
                uom,
                source_lookup,
            )
        )
        digest = hashlib.sha256(identity.casefold().encode()).hexdigest()[:24]
        suffix = cls._gas_suffix(ghg_unit)
        return f"legacy_{digest}_{suffix}"

    @staticmethod
    def _gas_suffix(ghg_unit: str) -> str:
        lowered = ghg_unit.casefold()
        if lowered == "kg co2e":
            return "1"
        if "co2" in lowered and "co2e" not in lowered:
            return "2"
        if "of co2" in lowered:
            return "2"
        if "ch4" in lowered:
            return "3"
        if "n2o" in lowered:
            return "4"
        return "5"

    @staticmethod
    def _normalize_ghg_unit(value: str) -> str:
        normalized = " ".join(value.split())
        lowered = normalized.casefold()
        if lowered.startswith("kg co2e of co2"):
            return "kg CO2e of CO2 per unit"
        if lowered.startswith("kg co2e of ch4"):
            return "kg CO2e of CH4 per unit"
        if lowered.startswith("kg co2e of n2o"):
            return "kg CO2e of N2O per unit"
        return normalized

    @staticmethod
    def _normalize_uom(value: str | None) -> str | None:
        if value is None:
            return None
        if value.casefold() == "litres":
            return "litres"
        return value

    @staticmethod
    def _text(value: object) -> str | None:
        if value is None:
            return None
        result = str(value).strip()
        return result or None

    @classmethod
    def _required_text(cls, value: object, field: str, row: int) -> str:
        result = cls._text(value)
        if result is None:
            raise PermanentIngestionError(f"DEFRA {field} is missing at row {row}")
        return result

    @staticmethod
    def _decimal(value: object, row: int) -> Decimal | None:
        if value is None or value == "":
            return None
        if isinstance(value, str) and (
            value.strip().startswith("<")
            or value.strip().casefold() in {"n/a", "na", "no data", "-"}
        ):
            return None
        try:
            return Decimal(str(value))
        except (InvalidOperation, ValueError) as error:
            raise PermanentIngestionError(f"invalid DEFRA value at row {row}: {value}") from error

    @staticmethod
    def _value_qualifier(value: object) -> str | None:
        if not isinstance(value, str):
            return None
        normalized = " ".join(value.split())
        if not normalized:
            return None
        try:
            Decimal(normalized)
        except InvalidOperation:
            return normalized
        return None
