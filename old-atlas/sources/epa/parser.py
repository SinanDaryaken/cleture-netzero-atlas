from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from atlas.domain.enums import FactorIntendedUse, FactorValueKind, GeographicFitType
from atlas.domain.models import ParsedRecord
from atlas.ingestion.errors import PermanentIngestionError

SHEET_NAME = "Emission Factors Hub"
EXPECTED_TABLE_FACTOR_COUNTS = {
    1: 121,
    2: 10,
    3: 115,
    4: 34,
    5: 40,
    6: 56,
    7: 1,
    8: 7,
    9: 183,
    10: 12,
    11: 32,
    12: 30,
}


class EpaRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    table_number: int = Field(ge=1, le=12)
    table_title: str
    source_key: str
    name: str
    category: str
    activity_type: str
    denominator_unit: str
    variant: str
    source_row: int
    co2: Decimal | None = None
    co2_unit: str | None = None
    ch4: Decimal | None = None
    ch4_unit: str | None = None
    n2o: Decimal | None = None
    n2o_unit: str | None = None
    direct_value: Decimal | None = None
    direct_unit: str | None = None
    factor_value_kind: FactorValueKind = FactorValueKind.CO2E_TOTAL
    intended_use: FactorIntendedUse = FactorIntendedUse.INVENTORY
    geographic_fit_type: GeographicFitType = GeographicFitType.COUNTRY_SPECIFIC
    origin_code: str = "US"
    applicable_codes: tuple[str, ...] = ("US",)
    reference_year: int | None = None
    scope: str | None = None
    lifecycle_stage: str | None = None
    system_boundary: str | None = None
    gwp_standard: str | None = "AR5"
    description: str | None = None
    attributes: dict[str, Any] = Field(default_factory=dict)
    original_unit: str

    def as_parsed_record(self) -> ParsedRecord:
        return ParsedRecord(
            data=self.model_dump(mode="json"),
            source_row=self.source_row,
            source_sheet=SHEET_NAME,
            source_table=f"Table {self.table_number}: {self.table_title}",
        )


class EpaParser:
    """Parser for the 12 logical tables in EPA's annual Hub workbook."""

    def __init__(self) -> None:
        self.metrics: dict[str, int] = {}

    def parse(self, path: Path) -> tuple[EpaRow, ...]:
        # The parser performs bounded random cell access across 12 tables. Normal
        # mode avoids openpyxl's expensive repeated streaming scans for cell().
        workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise PermanentIngestionError(f"EPA sheet {SHEET_NAME!r} is missing")
            sheet = workbook[SHEET_NAME]
            anchors = self._table_anchors(sheet)
            if set(anchors) != set(range(1, 13)):
                missing = sorted(set(range(1, 13)) - set(anchors))
                raise PermanentIngestionError(f"EPA workbook tables are missing: {missing}")

            rows: list[EpaRow] = []
            parsers = (
                self._table_1,
                self._table_2,
                self._table_3,
                self._table_4,
                self._table_5,
                self._table_6,
                self._table_7,
                self._table_8,
                self._table_9,
                self._table_10,
                self._table_11,
                self._table_12,
            )
            self.metrics = {"tables_detected": len(anchors), "na_values": 0}
            for number, parser in enumerate(parsers, start=1):
                parsed = parser(sheet, anchors[number])
                rows.extend(parsed)
                self.metrics[f"table_{number}_factors"] = len(parsed)
            if not rows:
                raise PermanentIngestionError("EPA workbook contains no factor rows")
            return tuple(rows)
        finally:
            workbook.close()

    @staticmethod
    def _table_anchors(sheet: Any) -> dict[int, tuple[int, str]]:
        anchors: dict[int, tuple[int, str]] = {}
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            marker = EpaParser._text(row[1] if len(row) > 1 else None)
            if marker is None or not marker.startswith("Table "):
                continue
            try:
                number = int(marker.removeprefix("Table ").strip())
            except ValueError:
                continue
            title = EpaParser._text(row[2] if len(row) > 2 else None) or marker
            anchors[number] = (index, title)
        return anchors

    def _table_1(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        rows: list[EpaRow] = []
        category = "Stationary combustion"
        source_unit: str | None = None
        for index in range(start + 3, sheet.max_row + 1):
            name = self._text(sheet.cell(index, 3).value)
            heat_header = self._text(sheet.cell(index, 4).value)
            if name == "Source:":
                break
            if heat_header and heat_header.lower().startswith("mmbtu per "):
                source_unit = heat_header.split(" per ", 1)[1]
                continue
            co2 = self._decimal(sheet.cell(index, 5).value)
            ch4 = self._decimal(sheet.cell(index, 6).value)
            n2o = self._decimal(sheet.cell(index, 7).value)
            if co2 is None and ch4 is None and n2o is None:
                if name:
                    category = name
                continue
            if not name:
                continue
            rows.append(
                self._component_row(
                    1,
                    title,
                    name,
                    category,
                    "fuel",
                    "mmBtu",
                    "heat_content",
                    index,
                    co2,
                    ch4,
                    n2o,
                    original_unit="kg CO2, g CH4, g N2O per mmBtu",
                    scope="scope_1",
                    boundary="combustion",
                )
            )
            source_co2 = self._decimal(sheet.cell(index, 8).value)
            source_ch4 = self._decimal(sheet.cell(index, 9).value)
            source_n2o = self._decimal(sheet.cell(index, 10).value)
            if source_co2 is not None and source_unit is not None:
                rows.append(
                    self._component_row(
                        1,
                        title,
                        name,
                        category,
                        "fuel",
                        source_unit,
                        "source_unit",
                        index,
                        source_co2,
                        source_ch4,
                        source_n2o,
                        original_unit=f"kg CO2, g CH4, g N2O per {source_unit}",
                        scope="scope_1",
                        boundary="combustion",
                    )
                )
        return rows

    def _table_2(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        rows: list[EpaRow] = []
        for index in range(start + 3, sheet.max_row + 1):
            name = self._text(sheet.cell(index, 3).value)
            if name == "Source:":
                break
            value = self._decimal(sheet.cell(index, 4).value)
            unit = self._text(sheet.cell(index, 5).value)
            if name and value is not None and unit:
                rows.append(
                    EpaRow(
                        table_number=2,
                        table_title=title,
                        source_key=name,
                        name=name,
                        category="Mobile combustion CO2",
                        activity_type="fuel",
                        denominator_unit=unit,
                        variant="co2",
                        source_row=index,
                        co2=value,
                        co2_unit="kg",
                        factor_value_kind=FactorValueKind.CO2_ONLY,
                        reference_year=None,
                        scope="scope_1",
                        system_boundary="tank_to_wheel",
                        original_unit=f"kg CO2/{unit}",
                    )
                )
        return rows

    def _table_3(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        return self._mobile_partial_rows(
            sheet,
            start + 3,
            title,
            3,
            vehicle_col=3,
            fuel_col=None,
            model_col=4,
            ch4_col=5,
            n2o_col=6,
            denominator="vehicle-mile",
        )

    def _table_4(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        return self._mobile_partial_rows(
            sheet,
            start + 3,
            title,
            4,
            vehicle_col=3,
            fuel_col=4,
            model_col=5,
            ch4_col=6,
            n2o_col=7,
            denominator="vehicle-mile",
        )

    def _table_5(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        return self._mobile_partial_rows(
            sheet,
            start + 3,
            title,
            5,
            vehicle_col=3,
            fuel_col=4,
            model_col=None,
            ch4_col=5,
            n2o_col=6,
            denominator="gallon",
        )

    def _mobile_partial_rows(
        self,
        sheet: Any,
        first_row: int,
        title: str,
        table_number: int,
        *,
        vehicle_col: int,
        fuel_col: int | None,
        model_col: int | None,
        ch4_col: int,
        n2o_col: int,
        denominator: str,
    ) -> list[EpaRow]:
        rows: list[EpaRow] = []
        vehicle: str | None = None
        fuel: str | None = None
        for index in range(first_row, sheet.max_row + 1):
            raw_vehicle = self._text(sheet.cell(index, vehicle_col).value)
            if raw_vehicle and raw_vehicle.startswith("Source:"):
                break
            vehicle = raw_vehicle or vehicle
            if fuel_col is not None:
                fuel = self._text(sheet.cell(index, fuel_col).value) or fuel
            model_year = self._text(sheet.cell(index, model_col).value) if model_col else None
            ch4 = self._decimal(sheet.cell(index, ch4_col).value)
            n2o = self._decimal(sheet.cell(index, n2o_col).value)
            if vehicle is None or (ch4 is None and n2o is None):
                continue
            parts = [vehicle]
            if fuel:
                parts.append(fuel)
            if model_year:
                parts.append(model_year)
            name = " / ".join(parts)
            rows.append(
                EpaRow(
                    table_number=table_number,
                    table_title=title,
                    source_key="|".join(parts),
                    name=name,
                    category="Mobile combustion non-CO2",
                    activity_type="transport",
                    denominator_unit=denominator,
                    variant="non_co2",
                    source_row=index,
                    ch4=ch4,
                    ch4_unit="g",
                    n2o=n2o,
                    n2o_unit="g",
                    factor_value_kind=FactorValueKind.NON_CO2_CO2E,
                    reference_year=2022,
                    scope="scope_1",
                    system_boundary="tank_to_wheel",
                    attributes={
                        "vehicle_type": vehicle,
                        "fuel_type": fuel,
                        "model_year": model_year,
                    },
                    original_unit=f"g CH4, g N2O/{denominator}",
                )
            )
        return rows

    def _table_6(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        rows: list[EpaRow] = []
        header = self._text(sheet.cell(start + 3, 3).value)
        first_name_candidate = self._text(sheet.cell(start + 5, 4).value)
        has_separate_grid_code = (header is not None and "acronym" in header.casefold()) or (
            first_name_candidate is not None and not self._is_decimal(first_name_candidate)
        )
        code_column = 3
        name_column = 4 if has_separate_grid_code else 3
        first_factor_column = 5 if has_separate_grid_code else 4
        for index in range(start + 5, sheet.max_row + 1):
            code = self._text(sheet.cell(index, code_column).value)
            if code and code.startswith("Source:"):
                break
            name = self._text(sheet.cell(index, name_column).value)
            if not code or not name:
                continue
            if not has_separate_grid_code:
                code = "US Average" if name.startswith("US Average") else name.split(" ", 1)[0]
            grid_loss = (
                self._decimal(sheet.cell(index, 11).value) if has_separate_grid_code else None
            )
            is_average = code == "US Average"
            for variant, column_offset, intended_use in (
                ("total_output", 0, FactorIntendedUse.INVENTORY),
                ("non_baseload", 3, FactorIntendedUse.AVOIDED_EMISSIONS),
            ):
                first_col = first_factor_column + column_offset
                rows.append(
                    EpaRow(
                        table_number=6,
                        table_title=title,
                        source_key=f"{code}|{variant}",
                        name=f"{name} / {variant.replace('_', ' ').title()}",
                        category="Electricity",
                        activity_type="electricity",
                        denominator_unit="MWh",
                        variant=variant,
                        source_row=index,
                        co2=self._required_decimal(sheet.cell(index, first_col).value, index),
                        co2_unit="lb",
                        ch4=self._required_decimal(sheet.cell(index, first_col + 1).value, index),
                        ch4_unit="lb",
                        n2o=self._required_decimal(sheet.cell(index, first_col + 2).value, index),
                        n2o_unit="lb",
                        intended_use=intended_use,
                        geographic_fit_type=(
                            GeographicFitType.COUNTRY_SPECIFIC
                            if is_average
                            else GeographicFitType.REGIONAL
                        ),
                        applicable_codes=("US",) if is_average else ("US", code),
                        reference_year=2023,
                        scope="scope_2",
                        system_boundary="generation",
                        attributes={
                            "grid_code": code,
                            "grid_name": name,
                            "grid_loss": grid_loss,
                        },
                        original_unit="lb CO2, lb CH4, lb N2O/MWh",
                    )
                )
        return rows

    def _table_7(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        index = start + 3
        name = self._required_text(sheet.cell(index, 3).value, index)
        return [
            self._component_row(
                7,
                title,
                name,
                "Purchased steam and heat",
                "energy",
                "mmBtu",
                "total",
                index,
                self._decimal(sheet.cell(index, 4).value),
                self._decimal(sheet.cell(index, 5).value),
                self._decimal(sheet.cell(index, 6).value),
                original_unit="kg CO2, g CH4, g N2O/mmBtu",
                scope="scope_2",
                boundary="combustion",
            )
        ]

    def _table_8(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        return self._distance_rows(sheet, start + 3, title, 8, end_marker="Source:")

    def _table_9(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        treatments = {
            4: "recycled",
            5: "landfilled",
            6: "combusted",
            7: "composted",
            8: "anaerobically_digested_dry",
            9: "anaerobically_digested_wet",
        }
        rows: list[EpaRow] = []
        for index in range(start + 4, sheet.max_row + 1):
            table_marker = self._text(sheet.cell(index, 2).value)
            if table_marker and table_marker.startswith("Table "):
                break
            material = self._text(sheet.cell(index, 3).value)
            if material and material.startswith("Source:"):
                break
            if not material:
                continue
            for column, treatment in treatments.items():
                raw_value = sheet.cell(index, column).value
                if self._text(raw_value) == "NA":
                    self.metrics["na_values"] = self.metrics.get("na_values", 0) + 1
                    continue
                value = self._decimal(raw_value)
                if value is None:
                    continue
                rows.append(
                    EpaRow(
                        table_number=9,
                        table_title=title,
                        source_key=f"{material}|{treatment}",
                        name=f"{material} / {treatment.replace('_', ' ').title()}",
                        category="Waste",
                        activity_type="waste",
                        denominator_unit="short ton",
                        variant=treatment,
                        source_row=index,
                        direct_value=value,
                        direct_unit="metric_tonne_co2e",
                        reference_year=2023,
                        scope="scope_3",
                        lifecycle_stage="end_of_life",
                        system_boundary="waste_treatment_including_transport_excluding_avoided",
                        gwp_standard="AR4",
                        attributes={"material": material, "treatment": treatment},
                        original_unit="metric tonne CO2e/short ton material",
                    )
                )
        return rows

    def _table_10(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        return self._distance_rows(sheet, start + 3, title, 10, end_marker="Source:")

    def _distance_rows(
        self, sheet: Any, first_row: int, title: str, table_number: int, *, end_marker: str
    ) -> list[EpaRow]:
        rows: list[EpaRow] = []
        for index in range(first_row, sheet.max_row + 1):
            name = self._text(sheet.cell(index, 3).value)
            if name and name.startswith(end_marker):
                break
            unit = self._text(sheet.cell(index, 7).value)
            if not name or not unit:
                continue
            reference_year = 2022
            origin_code = "US"
            geographic_fit_type = GeographicFitType.COUNTRY_SPECIFIC
            if table_number == 10:
                lowered = name.lower()
                if "intercity rail" in lowered:
                    reference_year = 2023
                elif "rail" in lowered:
                    reference_year = 2019
                elif "air travel" in lowered:
                    reference_year = 2018
                    origin_code = "GB"
                    geographic_fit_type = GeographicFitType.PROXY
            rows.append(
                self._component_row(
                    table_number,
                    title,
                    name,
                    "Freight transport" if table_number == 8 else "Business travel",
                    "transport",
                    unit,
                    "distance",
                    index,
                    self._decimal(sheet.cell(index, 4).value),
                    self._decimal(sheet.cell(index, 5).value),
                    self._decimal(sheet.cell(index, 6).value),
                    original_unit=f"kg CO2, g CH4, g N2O/{unit}",
                    scope="scope_3",
                    boundary="tank_to_wheel",
                    reference_year=reference_year,
                    origin_code=origin_code,
                    geographic_fit_type=geographic_fit_type,
                )
            )
        return rows

    def _table_11(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        rows: list[EpaRow] = []
        first_formula_candidate = self._text(sheet.cell(start + 3, 4).value)
        first_value_candidate = sheet.cell(start + 3, 5).value
        has_formula_column = "formula" in (
            self._text(sheet.cell(start + 2, 4).value) or ""
        ).casefold() or (
            first_formula_candidate is not None
            and not self._is_decimal(first_formula_candidate)
            and self._is_decimal(first_value_candidate)
        )
        for index in range(start + 3, sheet.max_row + 1):
            name = self._text(sheet.cell(index, 3).value)
            if name and name.startswith("Source:"):
                break
            formula = self._text(sheet.cell(index, 4).value) if has_formula_column else name
            value_column = 5 if has_formula_column else 4
            raw_value = sheet.cell(index, value_column).value
            if isinstance(raw_value, str) and raw_value.strip().startswith(">"):
                self.metrics["qualified_values"] = self.metrics.get("qualified_values", 0) + 1
                continue
            value = self._decimal(raw_value)
            if name and formula and value is not None:
                rows.append(self._characterization_row(11, title, name, formula, value, index))
        return rows

    def _table_12(self, sheet: Any, anchor: tuple[int, str]) -> list[EpaRow]:
        start, title = anchor
        rows: list[EpaRow] = []
        for index in range(start + 3, sheet.max_row + 1):
            name = self._text(sheet.cell(index, 3).value)
            if name and name.startswith("Source:"):
                break
            value = self._decimal(sheet.cell(index, 4).value)
            composition = self._text(sheet.cell(index, 5).value)
            if name and value is not None:
                rows.append(
                    self._characterization_row(
                        12, title, name, name, value, index, description=composition
                    )
                )
        return rows

    @staticmethod
    def _characterization_row(
        table_number: int,
        title: str,
        name: str,
        source_key: str,
        value: Decimal,
        index: int,
        *,
        description: str | None = None,
    ) -> EpaRow:
        return EpaRow(
            table_number=table_number,
            table_title=title,
            source_key=source_key,
            name=name,
            description=description,
            category="Global warming potential",
            activity_type="greenhouse_gas",
            denominator_unit="kg",
            variant="gwp_100_year",
            source_row=index,
            direct_value=value,
            direct_unit="kg_co2e_per_kg",
            factor_value_kind=FactorValueKind.CHARACTERIZATION_FACTOR,
            intended_use=FactorIntendedUse.CHARACTERIZATION,
            geographic_fit_type=GeographicFitType.GLOBAL,
            origin_code="GLOBAL",
            applicable_codes=("GLOBAL",),
            reference_year=2013,
            lifecycle_stage="use_phase",
            system_boundary="characterization",
            original_unit="kg CO2e/kg gas",
        )

    @staticmethod
    def _component_row(
        table_number: int,
        title: str,
        name: str,
        category: str,
        activity_type: str,
        denominator: str,
        variant: str,
        index: int,
        co2: Decimal | None,
        ch4: Decimal | None,
        n2o: Decimal | None,
        *,
        original_unit: str,
        scope: str,
        boundary: str,
        reference_year: int | None = None,
        origin_code: str = "US",
        geographic_fit_type: GeographicFitType = GeographicFitType.COUNTRY_SPECIFIC,
    ) -> EpaRow:
        return EpaRow(
            table_number=table_number,
            table_title=title,
            source_key=f"{name}|{variant}",
            name=name,
            category=category,
            activity_type=activity_type,
            denominator_unit=denominator,
            variant=variant,
            source_row=index,
            co2=co2,
            co2_unit="kg",
            ch4=ch4,
            ch4_unit="g",
            n2o=n2o,
            n2o_unit="g",
            reference_year=reference_year,
            scope=scope,
            system_boundary=boundary,
            origin_code=origin_code,
            applicable_codes=("US",),
            geographic_fit_type=geographic_fit_type,
            original_unit=original_unit,
        )

    @staticmethod
    def _text(value: object) -> str | None:
        if value is None:
            return None
        result = " ".join(str(value).replace("\xa0", " ").split())
        return result or None

    @classmethod
    def _required_text(cls, value: object, row: int) -> str:
        result = cls._text(value)
        if result is None:
            raise PermanentIngestionError(f"EPA required text is missing at row {row}")
        return result

    @staticmethod
    def _decimal(value: object) -> Decimal | None:
        if value is None or value == "":
            return None
        if isinstance(value, str) and value.strip().upper() == "NA":
            return None
        try:
            return Decimal(str(value))
        except (InvalidOperation, ValueError) as error:
            raise PermanentIngestionError(f"invalid EPA numeric value: {value}") from error

    @staticmethod
    def _is_decimal(value: object) -> bool:
        try:
            Decimal(str(value))
        except (InvalidOperation, ValueError):
            return False
        return True

    @classmethod
    def _required_decimal(cls, value: object, row: int) -> Decimal:
        result = cls._decimal(value)
        if result is None:
            raise PermanentIngestionError(f"EPA required numeric value is missing at row {row}")
        return result
