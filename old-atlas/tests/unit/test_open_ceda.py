from __future__ import annotations

import os
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from atlas.domain.enums import GeographicFitType
from atlas.domain.models import RawAssetReference
from atlas.ingestion.errors import PermanentIngestionError
from sources.open_ceda.normalizer import OpenCedaNormalizer
from sources.open_ceda.parser import OpenCedaParser


def _add_parameter_sheets(workbook: Workbook, *, release_year: int, base_year: int) -> None:
    exchange = workbook.create_sheet("Exchange rates")
    exchange.cell(4, 1, "Country Code")
    exchange.cell(4, 2, "Country Name")
    exchange.cell(4, 3, "LCU (local currency unit)")
    exchange.cell(4, 4, release_year)
    exchange.append([])
    exchange.cell(5, 1, "TUR")
    exchange.cell(5, 2, "Türkiye")
    exchange.cell(5, 3, "Turkish Lira")
    exchange.cell(5, 4, 32.5)

    purchaser = workbook.create_sheet("Purchaser - producer conversion")
    purchaser.cell(4, 2, "Oilseed farming")
    purchaser.cell(4, 3, "Grain farming")
    purchaser.cell(5, 2, "1111A0")
    purchaser.cell(5, 3, "1111B0")
    purchaser.cell(6, 2, 0.7)
    purchaser.cell(6, 3, 0.8)

    price_index = workbook.create_sheet("Sector level Price Index")
    price_index.cell(4, 2, "Oilseed farming")
    price_index.cell(4, 3, "Grain farming")
    price_index.cell(5, 2, "1111A0")
    price_index.cell(5, 3, "1111B0")
    price_index.cell(6, 1, base_year)
    price_index.cell(6, 2, 100)
    price_index.cell(6, 3, 100)

    workbook.create_sheet("Metadata")


def _fixture_2025(path: Path) -> None:
    workbook = Workbook()
    cover = workbook.active
    assert isinstance(cover, Worksheet)
    cover.title = "Cover"
    cover.cell(9, 4, datetime(2025, 11, 11))
    cover.cell(10, 4, "CEDA 2025")
    cover.cell(11, 4, "CC BY-SA 4.0")

    factors = workbook.create_sheet("GHG_t_Raw")
    factors.cell(1, 2, 2023)
    factors.cell(2, 2, "Producer price")
    factors.cell(2, 4, "US Dollar")
    factors.cell(3, 4, "Oilseed farming")
    factors.cell(3, 5, "Grain farming")
    factors.cell(4, 4, "1111A0")
    factors.cell(4, 5, "1111B0")
    for row, code, name, values in (
        (5, "AFG", "Afghanistan", (0.8, 1.2)),
        (6, "ROW", "Rest of World", (0.5, 0.7)),
    ):
        factors.cell(row, 1, code)
        factors.cell(row, 2, name)
        factors.cell(row, 3, "kgCO2e/US Dollar")
        factors.cell(row, 4, values[0])
        factors.cell(row, 5, values[1])

    regional = workbook.create_sheet("Regional Average EFs")
    regional.cell(3, 3, "Oilseed farming")
    regional.cell(3, 4, "Grain farming")
    regional.cell(4, 3, "1111A0")
    regional.cell(4, 4, "1111B0")
    regional.cell(5, 1, "Western Asia")
    regional.cell(5, 2, "kgCO2e/US Dollar")
    regional.cell(5, 3, 0.6)
    regional.cell(5, 4, 0.9)

    mapping = workbook.create_sheet("Country to region mapping")
    mapping.append(["Country Code", "UN Subregion"])
    mapping.append(["TUR", "Western Asia"])
    _add_parameter_sheets(workbook, release_year=2025, base_year=2023)
    workbook.save(path)


def _fixture_2024(path: Path) -> None:
    workbook = Workbook()
    factors = workbook.active
    assert isinstance(factors, Worksheet)
    factors.title = "Open CEDA"
    factors.cell(12, 4, "May 22nd, 2025")
    factors.cell(13, 4, "CEDA 2024")
    factors.cell(14, 4, "CC BY-SA 4.0")
    factors.cell(24, 3, 2022)
    factors.cell(25, 3, "Producer price")
    factors.cell(26, 3, "US Dollar")
    factors.cell(27, 5, "Oilseed farming")
    factors.cell(27, 6, "Grain farming")
    factors.cell(28, 5, "1111A0")
    factors.cell(28, 6, "1111B0")
    for row, code, name, values in (
        (29, "AFG", "Afghanistan", (0.9, 1.3)),
        (30, "ROW", "Rest of World", (0.6, 0.8)),
    ):
        factors.cell(row, 2, code)
        factors.cell(row, 3, name)
        factors.cell(row, 4, "kgCO2e/US Dollar")
        factors.cell(row, 5, values[0])
        factors.cell(row, 6, values[1])
    _add_parameter_sheets(workbook, release_year=2024, base_year=2022)
    workbook.save(path)


def _raw(path: Path, release_year: int) -> RawAssetReference:
    return RawAssetReference(
        bucket="atlas-raw",
        object_key=f"sources/open-ceda/{release_year}/test.xlsx",
        sha256="c" * 64,
        filename=path.name,
        source_url=(
            "https://open-ceda.s3.us-west-2.amazonaws.com/data/"
            f"Open%20CEDA%20{release_year}%20by%20Watershed.xlsx"
        ),
        size_bytes=path.stat().st_size,
        downloaded_at=datetime.now(UTC),
    )


def test_open_ceda_2025_separates_factors_parameters_and_regional_fallback(
    tmp_path: Path,
) -> None:
    path = tmp_path / "open-ceda-2025.xlsx"
    _fixture_2025(path)
    parser = OpenCedaParser(expected_sector_count=2, expected_country_count=2)
    parsed = parser.parse(path, release_year=2025)
    factors, observations, metrics = OpenCedaNormalizer().normalize(
        parsed.factors,
        parsed.parameters,
        raw=_raw(path, 2025),
        revision="d" * 64,
        parser_version="0.2.0",
    )

    assert parsed.base_year == 2023
    assert len(factors) == 6
    assert len(observations) == 5
    assert metrics["country_modelled_factors"] == 2
    assert metrics["regional_fallback_factors"] == 2
    assert metrics["proxy_factors"] == 2

    afghanistan = next(
        factor
        for factor in factors
        if factor.origin_geography and factor.origin_geography.code == "AF"
    )
    assert afghanistan.reference_year == 2023
    assert afghanistan.activity_unit == "USD_2023_producer_price"
    assert afghanistan.provenance.column_number == 4
    assert afghanistan.provenance.column_name == "1111A0"
    assert afghanistan.geographic_fit_type == GeographicFitType.COUNTRY_MODELLED

    exchange_rate = next(
        observation
        for observation in observations
        if observation.attributes["parameter_type"] == "exchange_rate"
    )
    assert exchange_rate.provenance.column_number == 4
    assert exchange_rate.provenance.column_name == "TUR"

    regional = next(
        factor for factor in factors if factor.geographic_fit_type == GeographicFitType.REGIONAL
    )
    assert [item.code for item in regional.applicable_geographies] == ["TR"]
    proxy = next(
        factor for factor in factors if factor.geographic_fit_type == GeographicFitType.PROXY
    )
    assert proxy.applicable_geographies[0].code == "GLOBAL"


def test_open_ceda_2024_keeps_release_and_model_base_year_distinct(tmp_path: Path) -> None:
    path = tmp_path / "open-ceda-2024.xlsx"
    _fixture_2024(path)
    parsed = OpenCedaParser(expected_sector_count=2, expected_country_count=2).parse(
        path, release_year=2024
    )

    assert parsed.release_year == 2024
    assert parsed.base_year == 2022
    assert parsed.published_at.year == 2025
    assert len(parsed.factors) == 4
    assert len(parsed.parameters) == 5


def test_open_ceda_rejects_schema_drift(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    workbook.save(path)

    with pytest.raises(PermanentIngestionError, match="workbook sheets changed"):
        OpenCedaParser(expected_sector_count=2, expected_country_count=2).parse(
            path, release_year=2025
        )


@pytest.mark.skipif(
    not os.getenv("ATLAS_OPEN_CEDA_2024_XLSX"),
    reason="set ATLAS_OPEN_CEDA_2024_XLSX to run the official Open CEDA snapshot",
)
def test_official_open_ceda_2024_snapshot() -> None:
    path = Path(os.environ["ATLAS_OPEN_CEDA_2024_XLSX"])
    parser = OpenCedaParser()
    parsed = parser.parse(path, release_year=2024)

    assert parsed.base_year == 2022
    assert len(parsed.factors) == 59600
    assert len(parsed.parameters) == 4236


@pytest.mark.skipif(
    not os.getenv("ATLAS_OPEN_CEDA_2025_XLSX"),
    reason="set ATLAS_OPEN_CEDA_2025_XLSX to run the official Open CEDA snapshot",
)
def test_official_open_ceda_2025_snapshot() -> None:
    path = Path(os.environ["ATLAS_OPEN_CEDA_2025_XLSX"])
    parser = OpenCedaParser()
    parsed = parser.parse(path, release_year=2025)

    assert parsed.base_year == 2023
    assert len(parsed.factors) == 68000
    assert len(parsed.parameters) == 4784
    assert parser.metrics["regional_fallback_factors"] == 8400
