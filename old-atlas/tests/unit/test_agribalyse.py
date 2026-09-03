from __future__ import annotations

import hashlib
import io
import json
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook

from atlas.domain.enums import (
    EnvironmentalEntityType,
    FactorIntendedUse,
    GeographicFitType,
)
from atlas.domain.models import RawAssetReference
from atlas.validation import QualityEngine
from sources.agribalyse.normalizer import AgribalyseNormalizer
from sources.agribalyse.parser import AgribalyseParser


def _workbook(role: str) -> bytes:
    workbook = Workbook()
    notice = workbook.active
    notice.title = "Notice"
    notice["A1"] = "Version : Agribalyse v3.2; Septembre 2024"
    if role == "food":
        sheet = workbook.create_sheet("Synthese")
        sheet.cell(row=3, column=1, value="Code AGB")
        sheet.cell(row=4, column=1, value="11172")
        sheet.cell(row=4, column=3, value="food group")
        sheet.cell(row=4, column=4, value="food subgroup")
        sheet.cell(row=4, column=5, value="Produit alimentaire")
        sheet.cell(row=4, column=6, value="Food product")
        sheet.cell(row=4, column=7, value=2)
        sheet.cell(row=4, column=8, value=0)
        sheet.cell(row=4, column=9, value="Ambiant")
        sheet.cell(row=4, column=10, value="PACK PROXY")
        sheet.cell(row=4, column=11, value="No preparation")
        sheet.cell(row=4, column=12, value=2.24)
        sheet.cell(row=4, column=14, value=7.58)
    else:
        title = (
            "AGB 3.2 agricole conventionnel"
            if role == "conventional"
            else "AGB 3.2 agricole biologique"
        )
        sheet = workbook.create_sheet(title)
        sheet.cell(row=4, column=1, value=f"Produit {role}")
        sheet.cell(row=4, column=2, value=f"Product {role} at farm gate {{FR}} U")
        sheet.cell(row=4, column=3, value="Culture")
        sheet.cell(row=4, column=4, value="Agricultural\\Plant production")
        sheet.cell(row=4, column=6, value=-0.42 if role == "organic" else 0.25)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _bundle(path: Path) -> None:
    assets = []
    contents = {}
    for role in ("conventional", "organic", "food"):
        filename = f"agribalyse-3.2-{role}.xlsx"
        content = _workbook(role)
        contents[filename] = content
        assets.append(
            {
                "role": role,
                "filename": filename,
                "source_url": f"https://example.test/{filename}",
                "sha256": hashlib.sha256(content).hexdigest(),
                "size_bytes": len(content),
            }
        )
    manifest = json.dumps(
        {"source": "AGRIBALYSE", "dataset_version": "3.2", "assets": assets},
        sort_keys=True,
        separators=(",", ":"),
    ).encode()
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("bundle-manifest.json", manifest)
        for filename, content in contents.items():
            bundle.writestr(filename, content)


def test_agribalyse_bundle_preserves_lca_semantics(tmp_path: Path) -> None:
    path = tmp_path / "agribalyse.zip"
    _bundle(path)
    parser = AgribalyseParser()
    rows = parser.parse(path)

    assert len(rows) == 3
    assert parser.metrics == {
        "parsed_rows": 3,
        "conventional_rows": 1,
        "organic_rows": 1,
        "food_rows": 1,
        "negative_lca_results": 1,
    }
    raw = RawAssetReference(
        bucket="atlas-raw",
        object_key="sources/agribalyse/fixture/bundle.zip",
        sha256="a" * 64,
        filename="bundle.zip",
        source_url="https://agribalyse.ademe.fr/",
        mime_type="application/zip",
        size_bytes=path.stat().st_size,
        downloaded_at=datetime.now(UTC),
        metadata={"acquisition_bundle": True},
    )
    factors, metrics = AgribalyseNormalizer().normalize(
        rows,
        raw=raw,
        dataset_version="3.2",
        parser_version="0.1.0",
    )

    assert len(factors) == 3
    assert metrics["default_match_eligible"] == 0
    assert all(factor.entity_type == EnvironmentalEntityType.LCA_RESULT for factor in factors)
    assert all(factor.intended_use == FactorIntendedUse.CHARACTERIZATION for factor in factors)
    assert all(factor.default_match_eligible is False for factor in factors)
    assert all(len(factor.source_factor_id or "") <= 255 for factor in factors)
    assert all(factor.provenance.original_file.endswith(".xlsx") for factor in factors)
    food = next(factor for factor in factors if factor.activity_type == "food_product_lca")
    assert food.geographic_fit_type == GeographicFitType.COUNTRY_MODELLED
    assert food.data_quality == "2.24"
    negative = next(factor for factor in factors if factor.factor_value < 0)
    assert not any(
        finding.rule_code == "factor.negative_value"
        for finding in QualityEngine().validate((negative,)).findings
    )
