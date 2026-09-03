from __future__ import annotations

import os
from collections import Counter
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook

from atlas.domain.enums import FactorIntendedUse, FactorValueKind, GeographicFitType
from atlas.domain.models import RawAssetReference
from atlas.infrastructure.http import HttpAssetFetcher
from sources.epa.normalizer import EpaNormalizer
from sources.epa.parser import EXPECTED_TABLE_FACTOR_COUNTS, EpaParser


async def test_epa_release_resolver_targets_archived_year() -> None:
    html = """
    <html><body>
      <a href="/files/ghg-emission-factors-hub-2023.xlsx">
        ARCHIVED 2023 GHG Emission Factors Hub (xlsx)
      </a>
      <a href="/files/ghg-emission-factors-hub-2025.xlsx">
        2025 GHG Emission Factors Hub (xlsx)
      </a>
    </body></html>
    """
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, text=html))
    )

    asset_url, year = await HttpAssetFetcher(client).resolve_epa_release(
        "https://www.epa.gov/climateleadership/ghg-emission-factors-hub",
        reference_year=2023,
    )
    await client.aclose()

    assert year == 2023
    assert asset_url == "https://www.epa.gov/files/ghg-emission-factors-hub-2023.xlsx"


def _table(sheet: object, row: int, number: int, title: str) -> None:
    sheet.cell(row, 2, f"Table {number}")
    sheet.cell(row, 3, title)


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Emission Factors Hub"

    _table(sheet, 1, 1, "Stationary Combustion")
    sheet.cell(4, 4, "mmBtu per gallon")
    sheet.cell(5, 3, "Liquid fuels")
    for column, value in enumerate(("Diesel", None, 70, 1, 0.1, 7, 0.1, 0.01), start=3):
        sheet.cell(6, column, value)
    sheet.cell(7, 3, "Source:")

    _table(sheet, 10, 2, "Mobile Combustion CO2")
    for column, value in enumerate(("Diesel", 10, "gallon"), start=3):
        sheet.cell(13, column, value)
    sheet.cell(14, 3, "Source:")

    _table(sheet, 20, 3, "On-Road Gasoline")
    for column, value in enumerate(("Passenger car", 2022, 0.1, 0.01), start=3):
        sheet.cell(23, column, value)
    sheet.cell(24, 3, "Source:")

    _table(sheet, 30, 4, "Diesel Vehicles")
    for column, value in enumerate(("Passenger car", "Diesel", 2022, 0.1, 0.01), start=3):
        sheet.cell(33, column, value)
    sheet.cell(34, 3, "Source:")

    _table(sheet, 40, 5, "Non-Road Vehicles")
    for column, value in enumerate(("Boat", "Diesel", 0.1, 0.01), start=3):
        sheet.cell(43, column, value)
    sheet.cell(44, 3, "Source:")

    _table(sheet, 50, 6, "Electricity")
    for column, value in enumerate(("TEST", "Test Grid", 100, 1, 0.1, 200, 2, 0.2, 0.04), start=3):
        sheet.cell(55, column, value)
    sheet.cell(56, 3, "Source:")

    _table(sheet, 60, 7, "Steam and Heat")
    for column, value in enumerate(("Steam", 60, 1, 0.1), start=3):
        sheet.cell(63, column, value)

    _table(sheet, 70, 8, "Freight")
    for column, value in enumerate(("Truck", 1, 0.1, 0.01, "vehicle-mile"), start=3):
        sheet.cell(73, column, value)
    sheet.cell(74, 3, "Source:")

    _table(sheet, 80, 9, "Waste")
    for column, value in enumerate(("Paper", 0.1, 0.2, "NA", "NA", "NA", "NA"), start=3):
        sheet.cell(84, column, value)
    sheet.cell(85, 3, "Source:")

    _table(sheet, 90, 10, "Business Travel")
    for column, value in enumerate(
        ("Air Travel - Short Haul", 1, 0.1, 0.01, "passenger-mile"), start=3
    ):
        sheet.cell(93, column, value)
    sheet.cell(94, 3, "Source:")

    _table(sheet, 100, 11, "GWP")
    for column, value in enumerate(("Methane", "CH4", 28), start=3):
        sheet.cell(103, column, value)
    sheet.cell(104, 3, "Source:")

    _table(sheet, 110, 12, "Blended Refrigerants")
    for column, value in enumerate(("R-TEST", 100, "50% A, 50% B"), start=3):
        sheet.cell(113, column, value)
    sheet.cell(114, 3, "Source:")
    workbook.save(path)


def _raw(path: Path) -> RawAssetReference:
    return RawAssetReference(
        bucket="atlas-raw",
        object_key="sources/epa/2025/test/epa.xlsx",
        sha256="a" * 64,
        filename="epa.xlsx",
        source_url="https://example.test/epa.xlsx",
        size_bytes=path.stat().st_size,
        downloaded_at=datetime.now(UTC),
    )


def test_legacy_epa_grid_and_gwp_layout_are_versioned(tmp_path: Path) -> None:
    path = tmp_path / "epa-legacy.xlsx"
    _workbook(path)
    workbook = __import__("openpyxl").load_workbook(path)
    sheet = workbook["Emission Factors Hub"]
    sheet.cell(53, 3, "eGRID Subregion")
    sheet.cell(55, 3, "TEST (Test Grid)")
    for column, value in enumerate((100, 1, 0.1, 200, 2, 0.2), start=4):
        sheet.cell(55, column, value)
    sheet.cell(55, 10, None)
    sheet.cell(55, 11, None)
    sheet.cell(102, 3, "Gas")
    sheet.cell(102, 4, "100-Year GWP")
    sheet.cell(103, 3, "CH4")
    sheet.cell(103, 4, 25)
    sheet.cell(103, 5, None)
    workbook.save(path)

    parser = EpaParser()
    rows = parser.parse(path)

    grid = next(row for row in rows if row.table_number == 6)
    gwp = next(row for row in rows if row.table_number == 11)
    assert grid.attributes["grid_code"] == "TEST"
    assert grid.attributes["grid_name"] == "TEST (Test Grid)"
    assert gwp.direct_value == Decimal("25")


def test_all_epa_tables_preserve_semantics_and_na(tmp_path: Path) -> None:
    path = tmp_path / "epa.xlsx"
    _workbook(path)
    parser = EpaParser()
    rows = parser.parse(path)
    factors, metrics = EpaNormalizer(Path("sources/epa/mappings.yaml")).normalize(
        rows, raw=_raw(path), release_year=2025, parser_version="0.1.0"
    )

    assert len(rows) == len(factors) == 15
    assert parser.metrics["tables_detected"] == 12
    assert parser.metrics["na_values"] == 4
    assert metrics["excluded_rows"] == 0
    assert len({factor.source_factor_id for factor in factors}) == 15

    kinds = Counter((factor.factor_value_kind, factor.intended_use) for factor in factors)
    assert kinds[(FactorValueKind.NON_CO2_CO2E, FactorIntendedUse.INVENTORY)] == 3
    assert kinds[(FactorValueKind.CO2E_TOTAL, FactorIntendedUse.AVOIDED_EMISSIONS)] == 1
    assert kinds[(FactorValueKind.CHARACTERIZATION_FACTOR, FactorIntendedUse.CHARACTERIZATION)] == 2

    grid = next(
        factor
        for factor in factors
        if factor.methodology.details.get("grid_code") == "TEST"
        and factor.intended_use == FactorIntendedUse.INVENTORY
    )
    assert grid.geographic_fit_type == GeographicFitType.REGIONAL
    assert grid.geography_level.value == "grid"
    assert grid.geographic_specificity == 6
    assert [geography.code for geography in grid.applicable_geographies] == ["US", "TEST"]

    air = next(factor for factor in factors if factor.name == "Air Travel - Short Haul")
    assert air.origin_geography is not None and air.origin_geography.code == "GB"
    assert [geography.code for geography in air.applicable_geographies] == ["US"]
    assert air.geographic_fit_type == GeographicFitType.PROXY

    waste = [factor for factor in factors if factor.source_category == "Waste"]
    assert len(waste) == 2
    assert all(factor.factor_unit == "kgCO2e/kg" for factor in waste)


@pytest.mark.skipif(
    not os.getenv("ATLAS_EPA_WORKBOOK"),
    reason="set ATLAS_EPA_WORKBOOK to run the official EPA workbook snapshot",
)
def test_official_epa_workbook_snapshot() -> None:
    parser = EpaParser()
    rows = parser.parse(Path(os.environ["ATLAS_EPA_WORKBOOK"]))

    assert len(rows) == 641
    assert parser.metrics["na_values"] == 183
    assert {
        number: parser.metrics[f"table_{number}_factors"] for number in range(1, 13)
    } == EXPECTED_TABLE_FACTOR_COUNTS
