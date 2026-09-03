from __future__ import annotations

import os
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook

from atlas.domain.enums import FactorValueKind, GeographicFitType, ProvenanceRole
from atlas.domain.models import RawAssetReference
from atlas.infrastructure.http import HttpAssetFetcher
from atlas.ingestion.errors import PermanentIngestionError
from sources.aib.normalizer import AibNormalizer
from sources.aib.parser import (
    ATTRIBUTE_HEADERS,
    RESIDUAL_MIX_HEADERS,
    WASTE_HEADERS,
    AibParser,
)


def _fixture_workbook(path: Path) -> None:
    workbook = Workbook()
    residual = workbook.active
    residual.title = "Residual Mixes"
    residual.append(list(RESIDUAL_MIX_HEADERS))
    co2 = workbook.create_sheet("CO2")
    co2.append(list(ATTRIBUTE_HEADERS))
    waste = workbook.create_sheet("RW")
    waste.append(list(WASTE_HEADERS))

    records = (
        ("FR", Decimal("17.11"), Decimal("2.32"), Decimal("0.8194")),
        ("GB", Decimal("280.64"), Decimal("0.34"), Decimal("0.725")),
        ("AT", "NA", "NA", "NA"),
        ("IT", Decimal("427.78"), Decimal("0.22"), Decimal("0.782")),
    )
    for code, co2_value, waste_value, untracked in records:
        shares: list[object] = [Decimal("0.2"), *(Decimal("0") for _ in range(6))]
        shares.extend([Decimal("0.3"), Decimal("0.5"), *(Decimal("0") for _ in range(5))])
        if code == "IT":
            shares[0] = "20.00%"
            residual_waste: object = waste_value
        elif code == "AT":
            shares = ["NA"] * 14
            residual_waste = "NA"
        else:
            residual_waste = waste_value
        residual.append([code, *shares, untracked, co2_value, residual_waste])
        co2.append([code, Decimal("1"), co2_value, Decimal("2")])
        waste.append([code, Decimal("1"), waste_value, Decimal("2")])
    workbook.save(path)


def _raw(path: Path) -> RawAssetReference:
    return RawAssetReference(
        bucket="atlas-raw",
        object_key="sources/aib/2025-v1.0/test.xlsx",
        sha256="a" * 64,
        filename=path.name,
        source_url="https://www.aib-net.org/residual-mix-2025.xlsx",
        size_bytes=path.stat().st_size,
        downloaded_at=datetime.now(UTC),
    )


def test_aib_parser_and_normalizer_preserve_direct_co2_semantics(tmp_path: Path) -> None:
    path = tmp_path / "aib-residual-mix-2025.xlsx"
    _fixture_workbook(path)
    parser = AibParser()
    rows = parser.parse(path)
    factors, metrics = AibNormalizer(Path("sources/aib/mappings.yaml")).normalize(
        rows,
        raw=_raw(path),
        release_year=2025,
        release_version="1.0",
        parser_version="0.1.0",
    )

    assert len(rows) == 4
    assert parser.metrics == {
        "countries": 4,
        "available_residual_mix": 3,
        "full_disclosure_no_factor": 1,
    }
    assert len(factors) == 3
    assert metrics["excluded_full_disclosure"] == 1
    assert {factor.origin_geography.code for factor in factors if factor.origin_geography} == {
        "FR",
        "GB",
        "IT",
    }

    france = next(factor for factor in factors if factor.origin_geography.code == "FR")
    assert france.factor_value == Decimal("0.01711")
    assert france.factor_unit == "kgCO2/kWh"
    assert france.factor_value_kind == FactorValueKind.CO2_ONLY
    assert france.gases.co2 == france.factor_value
    assert france.gases.co2e is None
    assert france.default_match_eligible is False
    assert france.geographic_fit_type == GeographicFitType.COUNTRY_SPECIFIC
    assert france.provenance.role == ProvenanceRole.CO2
    assert france.provenance.sheet == "Residual Mixes"
    assert france.methodology.details["radioactive_waste_mg_per_kwh"] == Decimal("2.32")

    italy = next(factor for factor in factors if factor.origin_geography.code == "IT")
    assert italy.factor_value == Decimal("0.42778")
    assert italy.methodology.details["generation_shares"]["renewable_total"] == Decimal("0.2")
    assert italy.methodology.details["radioactive_waste_mg_per_kwh"] == Decimal("0.22")


def test_aib_parser_rejects_schema_drift(tmp_path: Path) -> None:
    path = tmp_path / "aib-drift.xlsx"
    _fixture_workbook(path)
    workbook = __import__("openpyxl").load_workbook(path)
    workbook["Residual Mixes"]["C1"] = "Unexpected"
    workbook.save(path)

    with pytest.raises(PermanentIngestionError, match="schema changed"):
        AibParser().parse(path)


async def test_aib_release_resolver_selects_latest_versioned_workbook() -> None:
    html = """
    <html><body>
      <p>Results for calendar year 2025</p>
      <p>Version 1.0, 2026-05-26</p>
      <a href="/files/2024_residual_mix.xlsx">2024 Residual Mix Excel</a>
      <a href="/files/2025_residual_mix.xlsx">2025 Residual Mix Excel</a>
    </body></html>
    """
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, text=html))
    )
    release = await HttpAssetFetcher(client).resolve_aib_release(
        "https://www.aib-net.org/facts/european-residual-mix"
    )
    await client.aclose()

    assert release.year == 2025
    assert release.version == "1.0"
    assert release.asset_url == "https://www.aib-net.org/files/2025_residual_mix.xlsx"
    assert release.published_on == datetime(2026, 5, 26, tzinfo=UTC)


async def test_aib_release_resolver_targets_year_when_workbook_name_has_no_year() -> None:
    html = """
    <html><body>
      <p>Version 1.0, 2024-05-30</p>
      <a href="/files/RM-calculation-results.xlsx">Download Excel</a>
    </body></html>
    """
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, text=html))
    )

    release = await HttpAssetFetcher(client).resolve_aib_release(
        "https://www.aib-net.org/facts/european-residual-mix/2023",
        reference_year=2023,
    )
    await client.aclose()

    assert release.year == 2023
    assert release.version == "1.0"
    assert release.asset_url == "https://www.aib-net.org/files/RM-calculation-results.xlsx"
    assert release.published_on == datetime(2024, 5, 30, tzinfo=UTC)


async def test_aib_release_without_machine_readable_asset_is_explicit() -> None:
    html = """
    <html><body>
      <p>Version 1.1, 2022-05-31</p>
      <a href="/files/residual-mix-2021.pdf">Download report</a>
    </body></html>
    """
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, text=html))
    )

    with pytest.raises(PermanentIngestionError, match=r"machine-readable workbook.*2021"):
        await HttpAssetFetcher(client).resolve_aib_release(
            "https://www.aib-net.org/facts/european-residual-mix/2021",
            reference_year=2021,
        )
    await client.aclose()


@pytest.mark.skipif(
    not os.getenv("ATLAS_AIB_XLSX"),
    reason="set ATLAS_AIB_XLSX to run the official AIB snapshot",
)
def test_official_aib_snapshot() -> None:
    path = Path(os.environ["ATLAS_AIB_XLSX"])
    parser = AibParser()
    rows = parser.parse(path)
    factors, metrics = AibNormalizer(Path("sources/aib/mappings.yaml")).normalize(
        rows,
        raw=_raw(path),
        release_year=2025,
        release_version="1.0",
        parser_version="0.1.0",
    )

    assert len(rows) == 34
    assert len(factors) == 31
    assert metrics["excluded_full_disclosure"] == 3
    assert {row.country_code for row in rows if not row.available} == {"AT", "CH", "NL"}
    italy = next(factor for factor in factors if factor.origin_geography.code == "IT")
    assert italy.factor_value == Decimal("0.42778")
